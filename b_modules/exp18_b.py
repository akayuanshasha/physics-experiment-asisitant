"""用分光计测三棱镜折射率实验模块
===================
一级大物光学实验 —— 用分光计测量三棱镜的折射率（最小偏向角法）

实验内容（对应《用分光计测三棱镜折射率B》实验指导）：
  a) 基础内容（测三棱镜顶角 A，3 次）：
     1. 棱镜 AC 面、AB 面依次正对望远镜，记左、右游标读数
        θ1、θ2 与 θ1′、θ2′，
        A = 180° − ½[(θ1−θ1′)+(θ2−θ2′)]（两次读数差即顶角的补角）。
  b) 提高内容（最小偏向角法测折射率，3 次）：
     2. 汞灯绿谱线（546.1 nm）处于最小偏向角位置时记左、右游标
        θ1、θ2；取下棱镜（载物台保持不动），望远镜对准平行光管
        确定入射光方向，记 θ1′、θ2′，
        δ_min = ½[(θ1−θ1′)+(θ2−θ2′)]。
     3. 将 Ā 与 δ̄_min 代入 n = sin[(A+δ_min)/2]/sin(A/2)，
        计算折射率及其标准不确定度（误差传递）。
  c) 进阶内容（选做，本模块不处理其数据）：
     4. 测汞灯其他谱线（435.8 nm、577.0/579.0 nm 等）的最小偏向角，
        用柯西公式 n(λ)=a+b/λ²+c/λ⁴ 求色散系数。

不确定度约定（与 exp18_a 一致）：角度读数仪器允差 1′、估读允差 0.5′，
C = √3、P = 0.95。单次测量的四个读数以 ±½ 权重进入 A 与 δ_min，
单次测量值的合成不确定度恰等于单个读数的不确定度，故 3 次测量值
直接以读数允差作为 B 类分量参与 analyse()。

读数检查（分光计读数规则）：同一位置左、右游标读数差与 180° 之差应不超 3′；
两游标测得的转角相差过大（>3′）说明读数有误或刻度盘偏心未调好。
"""

from collections import Counter

from head import *  # 导入万能头
from docx.shared import Inches
from structured_support import (
    copied_tables, formatted, make_schema, make_table,
)
from theory_content import get_formulas, get_variables, get_table_theory
from sample_data_loader import load_sample_data
from optics_common import OpticsInputError, circular_difference, parse_angle

# 分光计游标读数允差（沿用原模块约定）：仪器允差 1′、估读允差 0.5′
_DELTA_INSTRUMENT = 1.0 / 60.0
_DELTA_ESTIMATE = 0.5 / 60.0

# 汞灯绿线波长（nm）：指导书给定，标注折射率结果对应的谱线
_LAMBDA_GREEN = 546.1

# 读数规则：两游标读数差与 180° 之差、两游标转角差的允许上限（分）
_VERNIER_LIMIT_ARCMIN = 3.0

# 顶角 A 常见范围（°）：指导书提示一般 50°~70°（常见 60°）
_A_RANGE = (40.0, 80.0)
# 最小偏向角常见范围（°）：指导书提示一般 30°~45°
_DELTA_RANGE = (15.0, 60.0)

# 每张表的四个读数列：c0/c1 为第一位置左、右游标，c2/c3 为第二位置左、右游标，
# c4 为只读计算列（顶角 A 或最小偏向角 δ_min）
_READING_COLS = ("c0", "c1", "c2", "c3")
_RESULT_COL = "c4"

_TABLE1_LABELS = [
    "AC面 θ1/(°)", "AC面 θ2/(°)", "AB面 θ1′/(°)", "AB面 θ2′/(°)", "顶角 A/(°)",
]
_TABLE2_LABELS = [
    "最小偏向位置 θ1/(°)", "最小偏向位置 θ2/(°)",
    "入射光位置 θ1′/(°)", "入射光位置 θ2′/(°)", "δ_min/(°)",
]


def name():
    return "用分光计测三棱镜折射率"


# ─────────────────────────────────────────────
# 数据解析辅助函数
# ─────────────────────────────────────────────
def _reading(row, col):
    """解析单个游标读数（支持 "123°30′" 或十进制度），空/非法返回 None。"""
    raw = row.get(col, "")
    if raw is None or str(raw).strip() == "":
        return None
    try:
        return parse_angle(raw, "游标读数")
    except OpticsInputError:
        return None


def _vernier_pair_difference(row, col_first, col_second):
    """一对左右游标读数间的圆周角差（度，取最小差）；任一缺失返回 None。"""
    first = _reading(row, col_first)
    second = _reading(row, col_second)
    if first is None or second is None:
        return None
    return circular_difference(first, second)


def _apex_angle(row):
    """该行顶角：A = 180° − ½[(θ1−θ1′)+(θ2−θ2′)]。

    左右游标两对读数各自的圆周差为载物台转角（顶角的补角），
    两读数取平均以消除偏心误差。任一读数缺失返回 None。
    """
    d1 = _vernier_pair_difference(row, "c0", "c2")
    d2 = _vernier_pair_difference(row, "c1", "c3")
    if d1 is None or d2 is None:
        return None
    return 180.0 - (d1 + d2) / 2.0


def _delta_min(row):
    """该行最小偏向角：δ_min = ½[(θ1−θ1′)+(θ2−θ2′)]。

    最小偏向位置与入射光位置各取圆周差，两游标平均。任一读数缺失返回 None。
    """
    d1 = _vernier_pair_difference(row, "c0", "c2")
    d2 = _vernier_pair_difference(row, "c1", "c3")
    if d1 is None or d2 is None:
        return None
    return (d1 + d2) / 2.0


def _row_has_data(row):
    """该行的四个读数列中是否有内容。"""
    return any(str(row.get(col, "")).strip() for col in _READING_COLS)


def _vernier_checks(row):
    """行级读数检查：同一位置左、右游标读数差与 180° 之差应不超 3′；
    两游标测得的转角相差过大（>3′）说明读数有误或偏心未调好。
    返回该行的问题描述列表（无问题为空）。"""
    issues = []
    for col_first, col_second in (("c0", "c1"), ("c2", "c3")):
        first = _reading(row, col_first)
        second = _reading(row, col_second)
        if first is None or second is None:
            continue
        offset = abs(circular_difference(first, second) - 180.0) * 60.0
        if offset > _VERNIER_LIMIT_ARCMIN:
            issues.append("两游标读数差与 180° 相差 {:.1f}′".format(offset))
    d1 = _vernier_pair_difference(row, "c0", "c2")
    d2 = _vernier_pair_difference(row, "c1", "c3")
    if d1 is not None and d2 is not None:
        gap = abs(d1 - d2) * 60.0
        if gap > _VERNIER_LIMIT_ARCMIN:
            issues.append("两游标测得转角相差 {:.1f}′".format(gap))
    return issues


def _font():
    """返回思源黑体 FontProperties，找不到则返回 None（用默认字体）。"""
    path = os.environ.get('_B_FONT_PATH')
    if path and os.path.exists(path):
        return matplotlib.font_manager.FontProperties(fname=path)
    return None


def _decorate(ax, title, xlabel, ylabel):
    zhfont = _font()
    if zhfont:
        ax.set_title(title, fontproperties=zhfont, fontsize=14)
        ax.set_xlabel(xlabel, fontproperties=zhfont, fontsize=12)
        ax.set_ylabel(ylabel, fontproperties=zhfont, fontsize=12)
        ax.legend(prop=zhfont, fontsize=10)
    else:
        ax.set_title(title, fontsize=14)
        ax.set_xlabel(xlabel, fontsize=12)
        ax.set_ylabel(ylabel, fontsize=12)
        ax.legend(fontsize=10)


def _add_doc_table(docu, headers, rows):
    table = docu.add_table(rows=1, cols=len(headers), style='Table Grid')
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for row in rows:
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
    return table


def _measurement_chart(workpath, values, mean, title, filename, fmt='o', color='#4472C4'):
    """某量 3 次测量值散点图（+ 平均值虚线）。"""
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(np.arange(1, len(values) + 1), values, fmt, color=color,
            markersize=7, label='各次测量', zorder=5)
    ax.axhline(mean, color='#3A9D5D', linestyle='--', linewidth=1.5,
               label='平均值 {:.4f}°'.format(mean))
    _decorate(ax, title, "测量序号", "角度 / °")
    ax.grid(True, linestyle='--', alpha=0.4)
    fig.tight_layout()
    fig.savefig(os.path.join(workpath, filename), dpi=150, bbox_inches='tight')
    plt.close(fig)
    return filename


# ─────────────────────────────────────────────
# 计算：顶角/最小偏向角统计 → 折射率
# ─────────────────────────────────────────────
def _calc_table(rows, kind):
    """通用计算：逐行算顶角 A（kind="apex"）或最小偏向角 δ_min（kind="delta"），
    行级读数检查 + 统计 + 不确定度 + 范围检查。返回富结果 dict。"""
    is_apex = kind == "apex"
    compute = _apex_angle if is_apex else _delta_min
    name_cn = "顶角 A" if is_apex else "最小偏向角 δ_min"
    avg_label = "平均顶角 Ā" if is_apex else "平均最小偏向角 δ̄_min"
    symbol = "A" if is_apex else "δ"
    range_lim = _A_RANGE if is_apex else _DELTA_RANGE

    values = []
    incomplete = 0
    issues = []
    for row in rows:
        if not _row_has_data(row):
            continue
        issues.extend(_vernier_checks(row))
        value = compute(row)
        if value is None:
            incomplete += 1
        else:
            values.append(value)

    warnings = []
    lines = []
    if incomplete:
        warnings.append("{} 行缺少部分游标读数，未计入统计。".format(incomplete))
    if not values:
        lines.append("暂无有效的{}数据。".format(name_cn))
        return {"values": values, "incomplete": incomplete, "res": None,
                "lines": lines, "warnings": warnings}

    if len(values) < 3:
        warnings.append("{}有效测量仅 {} 次，指导书建议 3 次。".format(name_cn, len(values)))
    lines.append("各次测得的{}（°）：{}".format(name_cn, '、'.join('%.5g' % v for v in values)))
    res = analyse(pd.Series(values), _DELTA_INSTRUMENT, _DELTA_ESTIMATE,
                  symbol, "°", confidence_C=3 ** 0.5)
    lines.append("{} ≈ {}°（U ≈ {}°，P=0.95）".format(
        avg_label, formatted(res.average, 3), formatted(res.unc, 3)))
    avg = float(res.average)
    if not (range_lim[0] <= avg <= range_lim[1]):
        warnings.append("{}平均值 {}° 超出常见范围（{}~{}°），请核对读数。".format(
            name_cn, formatted(avg, 3), range_lim[0], range_lim[1]))

    for text, count in Counter(issues).items():
        suffix = "（{} 行）".format(count) if count > 1 else ""
        warnings.append("读数检查：{}{} > 3′（分光计读数规则：两游标读数差与 180° 之差应不超 3′）。".format(
            text, suffix))

    return {"values": values, "incomplete": incomplete, "res": res,
            "lines": lines, "warnings": warnings}


def _calc_n(res_A, res_d):
    """折射率：n = sin[(Ā+δ̄min)/2]/sin(Ā/2)（角度以度代入，表达式内显式换算弧度）。"""
    res_n = analyse_com(
        "n=sin((A+δ)*pi/360)/sin(A*pi/360)",
        (("A", float(res_A.average), float(res_A.unc)),
         ("δ", float(res_d.average), float(res_d.unc))),
        (), "")
    n_value = float(res_n.ans)
    n_unc = float(res_n.unc)
    lines = [
        "顶角 Ā ≈ {}°、最小偏向角 δ̄_min ≈ {}°，代入最小偏向角公式".format(
            formatted(float(res_A.average), 3), formatted(float(res_d.average), 3)),
        "折射率 n ≈ {} ± {}（绿谱线 {} nm，P=0.95）".format(
            formatted(n_value, 4), formatted(n_unc, 4), _LAMBDA_GREEN),
    ]
    return {"res": res_n, "n": n_value, "unc": n_unc, "lines": lines}


# ─────────────────────────────────────────────
# 新版结构化接口
# ─────────────────────────────────────────────
def schema():
    sample_apex = load_sample_data("exp18_b", "顶角测量")
    sample_delta = load_sample_data("exp18_b", "最小偏向角测量")
    table1 = make_table(
        "table1", "三棱镜顶角 A 测量数据",
        _TABLE1_LABELS,
        sample=sample_apex,
        readonly=(4,),
        min_rows=3,
        description="θ1/θ2 为 AC 面正对望远镜时左、右游标读数，"
                    "θ1′/θ2′ 为 AB 面正对望远镜时的读数（各 3 次）。"
                    "角度可输入 \"123°30′\" 或十进制度（123.5 = 123°30′）。"
                    "顶角 A = 180° − ½[(θ1−θ1′)+(θ2−θ2′)] 由系统自动计算（已处理 0°/360° 跨界）。"
                    "读数规则：同一位置两游标读数差与 180° 之差应不超 3′。",
    )
    table1["calc"] = {"label": "计算顶角并统计"}
    table2 = make_table(
        "table2", "最小偏向角 δ_min 测量数据（绿谱线 546.1 nm）",
        _TABLE2_LABELS,
        sample=sample_delta,
        readonly=(4,),
        min_rows=3,
        description="θ1/θ2 为绿谱线处于最小偏向角位置时望远镜左、右游标读数，"
                    "θ1′/θ2′ 为取下棱镜后望远镜对准平行光管（入射光方向）的读数（各 3 次）。"
                    "δ_min = ½[(θ1−θ1′)+(θ2−θ2′)] 由系统自动计算（已处理 0°/360° 跨界），"
                    "两张表均有效时自动计算折射率。",
    )
    table2["calc"] = {"label": "计算最小偏向角与折射率"}

    return make_schema(
        (
            "最小偏向角法测三棱镜折射率：先以棱镜 AC、AB 面依次正对望远镜，"
            "记录左右游标读数（各 3 次）得顶角 " r"$A$" "；再测汞灯绿谱线（546.1 nm）的"
            "最小偏向角 " r"$\delta_{\min}$" "（最小偏向位置与入射光位置，各 3 次）；由 "
            r"$n = \frac{\sin[(\bar{A}+\bar{\delta}_{\min})/2]}{\sin(\bar{A}/2)}$"
            " 计算折射率及其不确定度。"
        ),
        [table1, table2],
        analysis_hints="顶角 A 一般为 50°~70°（常见 60°）；绿谱线最小偏向角 δ_min 一般为 30°~45°。"
                       "由 n = sin[(Ā+δ̄_min)/2]/sin(Ā/2) 得折射率，玻璃三棱镜典型值约 1.5~1.6"
                       "（K9 冕牌玻璃在 546.1 nm 处约 1.518）。"
                       "同一位置两游标读数差与 180° 之差应不超 3′，两游标测得转角也应一致。",
        preview_enabled=True,
        report_enabled=True,
        formulas=get_formulas("exp18_b"),
        variables=get_variables("exp18_b"),
        table_theory=get_table_theory("exp18_b"),
    )


def _compute_all(payload):
    """补全每行计算列（c4）并统计顶角/最小偏向角、计算折射率（富结果）。"""
    tables = copied_tables(payload)
    for row in tables.get("table1", []):
        value = _apex_angle(row)
        row["c4"] = formatted(value, 3) if value is not None else ""
    for row in tables.get("table2", []):
        value = _delta_min(row)
        row["c4"] = formatted(value, 3) if value is not None else ""

    calc_results = {}
    calc_messages = []
    fit_notes = {}
    apex_rows = tables.get("table1", [])
    delta_rows = tables.get("table2", [])

    if any(_row_has_data(r) for r in apex_rows):
        calc_results["table1"] = _calc_table(apex_rows, "apex")
        fit_notes["table1"] = [
            "顶角 A = 180° − ½[(θ1−θ1′)+(θ2−θ2′)]：两次读数之差是顶角的补角。",
            "左、右游标分别取圆周差再平均，消除刻度盘偏心误差。",
            "读数规则：同一位置两游标读数差与 180° 之差应不超 3′。",
        ]
    if any(_row_has_data(r) for r in delta_rows):
        calc_results["table2"] = _calc_table(delta_rows, "delta")
        fit_notes["table2"] = [
            "最小偏向角 δ_min = ½[(θ1−θ1′)+(θ2−θ2′)]：最小偏向位置与入射光方向之差。",
            "取下棱镜后望远镜对准平行光管，即入射光方向（载物台保持不动）。",
            "n = sin[(A+δ_min)/2]/sin(A/2)，需要顶角与最小偏向角两张表的有效数据。",
        ]

    apex = calc_results.get("table1", {})
    delta = calc_results.get("table2", {})
    if apex.get("res") is not None and delta.get("res") is not None:
        ncalc = _calc_n(apex["res"], delta["res"])
        delta["lines"].extend(ncalc["lines"])
        delta["res_n"] = ncalc
        delta["has_n"] = True
    elif apex.get("res") is not None:
        apex["lines"].append("填写表 2 的最小偏向角数据后可计算折射率。")

    return {"tables": tables, "calc_results": calc_results,
            "calc_messages": calc_messages, "fit_notes": fit_notes}


def preview(payload):
    """实时预计算接口：只返回前端需要且可 JSON 序列化的结论行。"""
    raw = _compute_all(payload)
    calc_results = {}
    for tid, v in raw["calc_results"].items():
        lines = list(v["lines"])
        lines.extend("注意：" + w for w in v["warnings"])
        calc_results[tid] = {"lines": lines}
    return {"tables": raw["tables"], "calc_results": calc_results,
            "calc_messages": raw["calc_messages"], "fit_notes": raw["fit_notes"]}


def handle_structured(workpath, payload):
    raw = _compute_all(payload)
    tables = raw["tables"]
    apex_rows = tables.get("table1", [])
    delta_rows = tables.get("table2", [])
    apex_calc = raw["calc_results"].get("table1")
    delta_calc = raw["calc_results"].get("table2")

    apex_data_rows = [
        [str(row.get(col, "")) for col in _READING_COLS + (_RESULT_COL,)]
        for row in apex_rows if _row_has_data(row)
    ]
    delta_data_rows = [
        [str(row.get(col, "")) for col in _READING_COLS + (_RESULT_COL,)]
        for row in delta_rows if _row_has_data(row)
    ]
    if not apex_data_rows and not delta_data_rows:
        return {"code": 1, "message": "请先填写测量数据表：至少需要顶角 A 或最小偏向角 δ_min 的游标读数。"}

    warnings = list(raw["calc_messages"])
    if apex_calc:
        warnings.extend(apex_calc["warnings"])
    if delta_calc:
        warnings.extend(delta_calc["warnings"])
    summary = []
    charts = []

    docu = Document()
    style_doc_font(docu)
    docu.add_heading(name(), level=0)

    # ── 一、三棱镜顶角 A 测量数据与统计 ─────────────────────────────────
    docu.add_heading("一、三棱镜顶角 A 测量数据与统计", level=1)
    if apex_data_rows:
        _add_doc_table(docu, _TABLE1_LABELS, apex_data_rows)
        docu.add_paragraph("θ1/θ2 为 AC 面正对望远镜时左、右游标读数，θ1′/θ2′ 为 AB 面正对望远镜时的读数；"
                           "顶角 A = 180° − ½[(θ1−θ1′)+(θ2−θ2′)] 由系统自动计算。")
    else:
        warnings.append("顶角 A 数据表为空，未进行顶角统计。")
    if apex_calc is not None and apex_calc.get("res") is not None:
        res_A = apex_calc["res"]
        insert_data(docu, "顶角 A", res_A, "word")
        summary.append("顶角 Ā ≈ {}°（U ≈ {}°，P=0.95）".format(
            formatted(float(res_A.average), 3), formatted(float(res_A.unc), 3)))
        fn = _measurement_chart(workpath, apex_calc["values"], float(res_A.average),
                                "三棱镜顶角三次测量", "chart_apex.png")
        charts.append({"filename": fn, "title": "三棱镜顶角三次测量"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
    elif apex_data_rows:
        warnings.append("顶角 A 无有效测量行，未进行顶角统计。")
    docu.add_paragraph()

    # ── 二、最小偏向角 δ_min 测量数据与统计 ─────────────────────────────
    docu.add_heading("二、最小偏向角 δ_min 测量数据与统计（绿谱线 {} nm）".format(_LAMBDA_GREEN), level=1)
    if delta_data_rows:
        _add_doc_table(docu, _TABLE2_LABELS, delta_data_rows)
        docu.add_paragraph("θ1/θ2 为绿谱线处于最小偏向角位置时望远镜左、右游标读数，"
                           "θ1′/θ2′ 为取下棱镜后望远镜对准平行光管（入射光方向）的读数；"
                           "δ_min = ½[(θ1−θ1′)+(θ2−θ2′)] 由系统自动计算。")
    else:
        warnings.append("最小偏向角 δ_min 数据表为空，未进行最小偏向角统计。")
    if delta_calc is not None and delta_calc.get("res") is not None:
        res_d = delta_calc["res"]
        insert_data(docu, "最小偏向角 δ_min", res_d, "word")
        summary.append("最小偏向角 δ̄_min ≈ {}°（U ≈ {}°，P=0.95）".format(
            formatted(float(res_d.average), 3), formatted(float(res_d.unc), 3)))
        fn = _measurement_chart(workpath, delta_calc["values"], float(res_d.average),
                                "绿谱线最小偏向角三次测量", "chart_delta.png",
                                fmt='s', color='#E74C3C')
        charts.append({"filename": fn, "title": "绿谱线最小偏向角三次测量"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
    elif delta_data_rows:
        warnings.append("最小偏向角 δ_min 无有效测量行，未进行最小偏向角统计。")
    docu.add_paragraph()

    # ── 三、三棱镜材料折射率 ────────────────────────────────────────────
    if delta_calc is not None and delta_calc.get("has_n"):
        ncalc = delta_calc["res_n"]
        docu.add_heading("三、三棱镜材料折射率", level=1)
        docu.add_paragraph("将顶角平均值 Ā 与绿谱线（{} nm）最小偏向角平均值 δ̄_min 代入最小偏向角公式"
                           "（角度以度代入，表达式内显式换算为弧度）：".format(_LAMBDA_GREEN))
        insert_data_com(docu, "折射率 n", ncalc["res"], "word")
        _add_doc_table(docu,
                       ["顶角 Ā/(°)", "最小偏向角 δ̄_min/(°)", "折射率 n"],
                       [[formatted(float(apex_calc["res"].average), 4),
                         formatted(float(delta_calc["res"].average), 4),
                         formatted(ncalc["n"], 4) + " ± " + formatted(ncalc["unc"], 4)]])
        docu.add_paragraph("三棱镜材料对绿谱线（{} nm）的折射率为 n ≈ {} ± {}（P = 0.95）。".format(
            _LAMBDA_GREEN, formatted(ncalc["n"], 4), formatted(ncalc["unc"], 4)))
        summary.append("折射率 n ≈ {} ± {}（绿谱线 {} nm，P=0.95）".format(
            formatted(ncalc["n"], 4), formatted(ncalc["unc"], 4), _LAMBDA_GREEN))
    else:
        warnings.append("折射率需要顶角 A 与最小偏向角 δ_min 各有至少 1 次有效测量，未计算折射率。")

    docu.save(os.path.join(workpath, name() + ".docx"))

    if not summary:
        summary.append("测量数据已记录，统计不完整部分请查看数据检查提示。")
    return {
        "code": 0,
        "summary": summary,
        "warnings": warnings,
        "charts": charts,
        "document": name() + ".docx",
    }


# ─────────────────────────────────────────────
# 旧版接口（向后兼容）
# ─────────────────────────────────────────────
def handle(workpath, extension):
    """旧版 CSV 单文件接口（兼容 b_adapter 等旧调用路径）：
    CSV 的 10 列依次对应 table1（4 读数 + A）与 table2（4 读数 + δ_min），
    A / δ_min 列可留空；还原为结构化数据后统一交由 handle_structured 处理。"""
    try:
        excelpath = workpath + name() + '.' + extension
        column_names = ["c{}".format(i) for i in range(10)]

        if extension == 'csv':
            with open(excelpath, 'rb') as f:
                encode = chardet.detect(f.read())["encoding"]
            data = pd.read_csv(excelpath, header=0, names=column_names, encoding=encode)
        else:
            data = pd.read_excel(excelpath, header=0, names=column_names)

        os.remove(excelpath)

        def _to_rows(start):
            rows = []
            for _, record in data.iterrows():
                row = {}
                for offset in range(4):
                    value = record["c{}".format(start + offset)]
                    row["c{}".format(offset)] = str(value).strip() if pd.notna(value) else ""
                rows.append(row)
            return [row for row in rows if _row_has_data(row)]

        payload = {"parameters": {}, "tables": {
            "table1": _to_rows(0),
            "table2": _to_rows(5),
        }}
        result = handle_structured(workpath, payload)
        return 0 if result.get("code") == 0 else 1
    except Exception:
        traceback.print_exc()  # 打印错误
        return 1  # 若失败，返回1
