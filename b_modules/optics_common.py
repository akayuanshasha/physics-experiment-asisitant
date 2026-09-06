"""二级光学实验共用的数据校验、绘图和 Word 报告工具。"""

from pathlib import Path
import re

from head import *
from docx.shared import Inches, Pt


class OpticsInputError(ValueError):
    """用户输入不满足实验要求。"""


def get_parameter(payload, key, default=None, required=False, cast=float):
    value = (payload.get("parameters") or {}).get(key, default)
    if value in (None, ""):
        if required:
            raise OpticsInputError(f"缺少参数：{key}")
        return default
    if cast is None:
        return value
    try:
        return cast(value)
    except (TypeError, ValueError) as exc:
        raise OpticsInputError(f"参数“{key}”格式不正确：{value}") from exc


def get_text(payload, key, default=""):
    value = (payload.get("text_fields") or {}).get(key, default)
    return str(value).strip() if value is not None else default


def get_rows(payload, table_id, required=False, min_rows=0):
    rows = (payload.get("tables") or {}).get(table_id, [])
    clean_rows = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        clean = {str(k): ("" if v is None else str(v).strip()) for k, v in row.items()}
        if any(value != "" for value in clean.values()):
            clean_rows.append(clean)
    if required and len(clean_rows) < max(1, min_rows):
        raise OpticsInputError(f"数据表“{table_id}”至少需要 {max(1, min_rows)} 行有效数据")
    if min_rows and clean_rows and len(clean_rows) < min_rows:
        raise OpticsInputError(f"数据表“{table_id}”至少需要 {min_rows} 行有效数据")
    return clean_rows


def numeric_column(rows, key, label=None, allow_empty=False):
    values = []
    for index, row in enumerate(rows, 1):
        raw = row.get(key, "")
        if raw == "" and allow_empty:
            values.append(np.nan)
            continue
        try:
            value = float(raw)
        except (TypeError, ValueError) as exc:
            raise OpticsInputError(
                f"{label or key}第 {index} 行不是有效数字：{raw or '空值'}"
            ) from exc
        if not np.isfinite(value):
            raise OpticsInputError(f"{label or key}第 {index} 行必须是有限数字")
        values.append(value)
    return np.asarray(values, dtype=float)


def create_document(title, subtitle=None):
    doc = Document()
    style_doc_font(doc)  # 统一标题/正文/表头为微软雅黑加粗
    doc.styles["Normal"].font.size = Pt(10.5)
    heading = doc.add_heading(title, level=0)
    heading.alignment = 1
    if subtitle:
        paragraph = doc.add_paragraph(subtitle)
        paragraph.alignment = 1
    return doc


def add_key_values(doc, title, values):
    doc.add_heading(title, level=1)
    for label, value in values:
        doc.add_paragraph(f"{label}：{value}")


def add_summary(doc, lines, title="计算结果"):
    doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(str(line), style="List Bullet")


def add_text_section(doc, title, content):
    if content:
        doc.add_heading(title, level=1)
        doc.add_paragraph(content)


def add_table(doc, title, headers, rows):
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers), style="Table Grid")
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True  # 表头加粗
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    return table


def configure_plotting():
    font_path = Path(__file__).resolve().parent.parent / "SourceHanSansSC-Regular.otf"
    if font_path.exists():
        matplotlib.font_manager.fontManager.addfont(str(font_path))
        plt.rcParams["font.family"] = "Source Han Sans SC"
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["xtick.direction"] = "in"
    plt.rcParams["ytick.direction"] = "in"


def save_figure(fig, workpath, filename):
    configure_plotting()
    path = Path(workpath) / filename
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return filename


def linear_fit(x, y):
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x = x[mask]
    y = y[mask]
    if len(x) < 2 or np.allclose(x, x[0]):
        raise OpticsInputError("线性拟合至少需要两个横坐标不同的有效数据点")
    result = scipy.stats.linregress(x, y)
    predicted = result.slope * x + result.intercept
    residuals = y - predicted
    return {
        "x": x,
        "y": y,
        "slope": float(result.slope),
        "intercept": float(result.intercept),
        "r": float(result.rvalue),
        "r2": float(result.rvalue ** 2),
        "slope_stderr": float(result.stderr or 0.0),
        "intercept_stderr": float(result.intercept_stderr or 0.0),
        "predicted": predicted,
        "residuals": residuals,
    }


_ANGLE_PATTERN = re.compile(
    r"^\s*([+-]?\d+(?:\.\d+)?)\s*(?:°|度|\s)\s*(\d+(?:\.\d+)?)?\s*(?:['′分])?\s*$"
)


def parse_angle(value, label="角度"):
    """接受十进制度或“123°30′ / 123 30”格式，返回十进制度。"""
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip().replace("º", "°")
        try:
            result = float(text)
        except ValueError:
            match = _ANGLE_PATTERN.match(text)
            if not match:
                raise OpticsInputError(f"{label}格式不正确：{value}")
            degrees = float(match.group(1))
            minutes = float(match.group(2) or 0.0)
            sign = -1.0 if degrees < 0 else 1.0
            result = degrees + sign * minutes / 60.0
    if not np.isfinite(result):
        raise OpticsInputError(f"{label}必须是有限数字")
    return result % 360.0


def circular_difference(angle_a, angle_b, period=360.0):
    """返回两个圆周角读数之间的最小绝对差。"""
    diff = abs((float(angle_a) - float(angle_b)) % period)
    return min(diff, period - diff)


def circular_signed_difference(angle_from, angle_to, period=360.0):
    return (float(angle_to) - float(angle_from) + period / 2) % period - period / 2


def add_charts_to_document(doc, workpath, charts):
    if not charts:
        return
    doc.add_heading("实验图表", level=1)
    for chart in charts:
        path = Path(workpath) / chart["filename"]
        if not path.exists():
            continue
        doc.add_paragraph(chart.get("title", path.stem))
        doc.add_picture(str(path), width=Inches(6.2))
        doc.paragraphs[-1].alignment = 1


def finish_report(doc, workpath, experiment_name, summary, charts=None, warnings=None):
    charts = charts or []
    warnings = warnings or []
    if warnings:
        add_summary(doc, warnings, title="数据检查提示")
    add_charts_to_document(doc, workpath, charts)
    output_path = Path(workpath) / f"{experiment_name}.docx"
    doc.save(output_path)
    return {
        "code": 0,
        "summary": [str(item) for item in summary],
        "warnings": [str(item) for item in warnings],
        "charts": charts,
    }


def error_result(exc):
    if isinstance(exc, OpticsInputError):
        return {"code": 1, "message": str(exc)}
    traceback.print_exc()
    return {"code": 1, "message": f"处理失败：{exc}"}
