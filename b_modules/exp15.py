"""硅光电池光电特性实验模块
==========================
二级大物电磁学实验 —— 硅光电池

实验内容（对应《硅光电池B》实验指导）：
1. 暗伏安特性测量（无光照正向偏压下 I-U 关系）
2. 输出特性测量（负载电阻 R 与工作电压 U，求 Isc/Uoc/Pm/Rm/FF）
3. 开路电压 Uoc、短路电流 Isc 与光照 L 特性测量
4. （进阶）输出电压 U 与光照 L 特性测量
5. （高阶）反向偏压下输出电压 U 与光照 L 特性测量

物理背景：
硅光电池基于 P-N 结光伏效应。
无光照时等效二极管：I = I0·(e^{qU/kT} - 1)
有光照时：I = I0·(e^{qU/kT} - 1) - Iph
短路电流 Isc = Iph ∝ L（与光照强度成正比）
开路电压 Uoc = (kT/q)·ln(1 + Isc/I0) ∝ ln L
输出功率 P = U·I，最大输出功率 Pm 对应最佳负载 Rm，
填充因子 FF = Pm / (Uoc·Isc)。
光照强度换算：d=50cm 时 L=40lx，故 L = 40·(50/d)²。

网页流程（新版结构化接口）：
1. schema() 声明五张数据表与实验参数；
2. preview() 实时给出暗伏安二极管拟合、输出特性参数（全部光照距离列）、
   Isc/Uoc-光照拟合与 U-L 拟合，并补全表 3 的短路电流列；
3. handle_structured() 生成各类特性曲线与 Word 数据处理文档。

旧版 CSV 流程 handle() 保留，供旧接口与 AI 助教调用。
"""

from head import *
from docx.shared import Inches
from structured_support import (
    as_number, copied_tables, formatted, make_schema, make_table,
)
from theory_content import get_formulas, get_variables, get_table_theory
from sample_data_loader import load_sample_data_numeric

_D50 = 50.0   # 标定距离 cm
_L50 = 40.0   # d=50cm 时对应的光照强度 lx

# 输出特性表的四个光照距离列：列 id、标签、灯距 cm
_OUTPUT_COLUMNS = (("c4", "d=50cm", 50.0), ("c3", "d=40cm", 40.0),
                   ("c2", "d=30cm", 30.0), ("c1", "d=20cm", 20.0))


def name():
    return "硅光电池"


def _light_lx(d):
    """灯距 d(cm) 换算为光照强度 L(lx)，L = 40·(50/d)²。"""
    d = float(d)
    if d <= 0:
        return None
    return _L50 * (_D50 / d) ** 2


def _font():
    """返回思源黑体 FontProperties，找不到则返回 None（用 matplotlib 默认字体）。"""
    path = os.environ.get('_B_FONT_PATH')
    if path and os.path.exists(path):
        return matplotlib.font_manager.FontProperties(fname=path)
    return None


def _has_value(v):
    """判断单元格是否填写了内容。"""
    if v is None:
        return False
    if isinstance(v, str):
        return bool(v.strip())
    return True


def _xy_pairs(rows, x_col, y_col):
    """从行对象中提取 (x, y) 数值对，跳过缺失项。"""
    xs, ys = [], []
    for r in rows:
        x = as_number(r.get(x_col))
        y = as_number(r.get(y_col))
        if x is not None and y is not None:
            xs.append(x)
            ys.append(y)
    return np.array(xs, dtype=float), np.array(ys, dtype=float)


def _fit_diode(U, I):
    """线性化拟合暗伏安特性 I = I0·(e^{bU}-1)，返回 (I0, b)。

    取 I>0 的点做 ln(I) 对 U 的线性拟合：ln I = ln I0 + b·U。
    """
    U = np.asarray(U, dtype=float)
    I = np.asarray(I, dtype=float)
    pos = I > 0
    if pos.sum() < 2:
        return None, None
    b, ln_a = np.polyfit(U[pos], np.log(I[pos]), 1)
    return float(np.exp(ln_a)), float(b)


def _log_fit(x, y):
    """对数拟合 y = A·ln(x) + B（要求 x>0），返回 (A, B)。"""
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    mask = x > 0
    if mask.sum() < 2:
        return None, None
    A, B = np.polyfit(np.log(x[mask]), y[mask], 1)
    return float(A), float(B)


def _log_eq(var, A, B):
    """格式化对数拟合方程「var = A·lnL + B」的文本，正确处理负号。"""
    sign = "-" if B < 0 else "+"
    return "{} = {}·lnL {} {}".format(var, formatted(A), sign, formatted(abs(B)))


def _linear_eq(var, k, b):
    """格式化线性拟合方程「var = k·L + b」的文本，正确处理负号。"""
    sign = "-" if b < 0 else "+"
    return "{} = {}·L {} {}".format(var, formatted(k), sign, formatted(abs(b)))


def _decorate(ax, title, xlabel, ylabel):
    zhfont = _font()
    if zhfont:
        ax.set_title(title, fontproperties=zhfont, fontsize=14)
        ax.set_xlabel(xlabel, fontproperties=zhfont, fontsize=12)
        ax.set_ylabel(ylabel, fontproperties=zhfont, fontsize=12)
        ax.legend(prop=zhfont, fontsize=10)
    else:
        ax.set_title(title, fontsize=14)
        ax.set_xlabel(xlabel, fontsize=12)
        ax.set_ylabel(ylabel, fontsize=12)
        ax.legend(fontsize=10)


def _save_chart(workpath, filename, title, xlabel, ylabel, x, y,
                fit_x=None, fit_y=None, fit_label=None, mark=None):
    """绘制散点 + 拟合曲线并保存到 workpath，返回文件名。"""
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(x, y, 'o', color='#4472C4', markersize=5, label='测量数据', zorder=5)
    if fit_x is not None and fit_y is not None:
        ax.plot(fit_x, fit_y, '-', color='#E74C3C', linewidth=2,
                label=fit_label or '拟合曲线')
    if mark is not None:
        ax.plot(mark[0], mark[1], 's', color='#2ECC71', markersize=9,
                label=mark[2] if len(mark) > 2 else '最佳工作点')
    _decorate(ax, title, xlabel, ylabel)
    ax.grid(True, linestyle='--', alpha=0.4)
    fig.tight_layout()
    path = os.path.join(workpath, filename)
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)
    return filename


def _output_pr_chart(workpath, series, primary_label):
    """全部光照距离列叠加的输出特性 P-R 曲线 + 主列（d=50cm）最大功率点标注。"""
    fig, ax = plt.subplots(figsize=(8, 5))
    colors = ("#4472C4", "#E74C3C", "#2ECC71", "#F39C12")
    for s, color in zip(series, colors):
        ax.plot(s["R"], s["P"] * 1000, 'o-', color=color, markersize=4,
                linewidth=1.2, label=s["label"])
    primary = next((s for s in series if s["label"] == primary_label), series[0])
    ax.plot([primary["rm"]], [primary["pm"] * 1000], 's', color="#8E44AD",
            markersize=10, label="最大功率点（{}）".format(primary["label"]))
    _decorate(ax, "硅光电池输出特性曲线 P-R", "R / Ω", "P / mW")
    ax.grid(True, linestyle='--', alpha=0.4)
    fig.tight_layout()
    path = os.path.join(workpath, "chart_output_pr.png")
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)
    return "chart_output_pr.png"


def _add_doc_table(docu, headers, rows):
    table = docu.add_table(rows=1, cols=len(headers), style='Table Grid')
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for row in rows:
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
    return table


# ── 各表实时计算（preview 与 handle_structured 共用）────────────────────


def _calc_dark(t1):
    """表 1 暗伏安特性：指数拟合 I = I0·(e^{bU}-1)，给出热电压与结温估计。"""
    U1, I1 = _xy_pairs(t1, "c0", "c1")
    if len(U1) < 2:
        raise ValueError("暗伏安特性至少需要 2 组有效数据（U 与 I 均填写）。")
    i0, b = _fit_diode(U1, I1)
    if i0 is None or b is None or b <= 0:
        raise ValueError("暗伏安特性拟合失败：需要至少 2 个正电流数据点。")
    kt_q = 1.0 / b          # V
    temp = 11604.5 / b      # K（k/q ≈ 8.617e-5 V/K 的倒数）
    lines = [
        "拟合方程：I = I₀·(e^{{bU}} − 1)，I₀ = {} mA，b = {} 1/V".format(
            formatted(i0), formatted(b)),
        "热电压 kT/q = {} mV，结温估计 T ≈ {} K".format(
            formatted(kt_q * 1000), round(temp)),
    ]
    if 250 <= temp <= 350:
        lines.append("结温估计接近室温（250~350 K），数据合理。")
    elif temp > 350:
        lines.append("注意：结温估计偏高（>350 K），请检查电压、电流数据。")
    else:
        lines.append("注意：结温估计偏低（<250 K），请检查电压、电流数据。")
    return {"lines": lines, "U": list(U1), "I": list(I1),
            "i0": i0, "b": b, "kt_q": kt_q, "temp": temp}


def _calc_output(t2):
    """表 2 输出特性：分析全部四个光照距离列，求 Uoc/Isc/Pm/Rm/FF。"""
    series = []
    notes = []
    lines = []
    for cid, lbl, d_cm in _OUTPUT_COLUMNS:
        R, U = _xy_pairs(t2, "c0", cid)
        if len(R) < 2:
            continue
        mask = R > 0
        if not mask.all():
            notes.append("{}：{} 个负载电阻非正数据点已跳过（电阻需 > 0）。".format(
                lbl, int((~mask).sum())))
            R, U = R[mask], U[mask]
        if len(R) < 2:
            continue
        I = U / R
        P = U * I
        imax = int(np.argmax(P))
        pm, rm, um, im = float(P[imax]), float(R[imax]), float(U[imax]), float(I[imax])
        uoc = float(U[np.argmax(R)])
        isc = float(I[np.argmin(R)])
        ff = pm / (uoc * isc) if uoc * isc > 0 else None
        L = _light_lx(d_cm)
        series.append({"label": lbl, "L": L, "R": R, "U": U, "I": I, "P": P,
                       "pm": pm, "rm": rm, "um": um, "im": im,
                       "uoc": uoc, "isc": isc, "ff": ff})
        lines.append("{}（L ≈ {:.1f} lx）：Uoc ≈ {} V，Isc ≈ {} mA，"
                     "Pm = {} mW（Rm = {} Ω，Um = {} V，Im = {} mA），FF ≈ {}".format(
                         lbl, L, formatted(uoc), formatted(isc * 1000),
                         formatted(pm * 1000), formatted(rm), formatted(um),
                         formatted(im * 1000),
                         formatted(ff) if ff is not None else "—"))
    if not series:
        raise ValueError("输出特性表至少需要 2 组有效数据（R 与某一距离列的 U）。")
    lines.append("注：此表 Uoc/Isc 为最大/最小负载下的近似值，可与表 3 实测值对照。")
    lines.extend(notes)
    return {"lines": lines, "series": series}


def _calc_uoc_isc(t3, r0):
    """表 3 开路电压、短路电流与光照：Isc=Ur/R₀，Isc-L 线性拟合、Uoc-L 对数拟合。"""
    if r0 is None or r0 <= 0:
        raise ValueError("参数「取样电阻 R₀」必须是正数，才能换算短路电流 Isc = Ur/R₀。")
    records = []
    for row in t3:
        d = as_number(row.get("c0"))
        uoc = as_number(row.get("c1"))
        ur = as_number(row.get("c2"))
        if d is None and uoc is None and ur is None:
            continue
        if d is None or d <= 0:
            raise ValueError("表 3 中灯距 d 需要为大于 0 的数值。")
        L = _light_lx(d)
        records.append({"d": d, "L": L, "uoc": uoc, "ur": ur,
                        "isc": ur / r0 if ur is not None else None})
    if not records:
        raise ValueError("开路电压、短路电流表没有有效数据行。")
    lines = []
    isc_list = [(x["L"], x["isc"]) for x in records if x["isc"] is not None]
    uoc_list = [(x["L"], x["uoc"]) for x in records if x["uoc"] is not None]
    isc_fit = None
    uoc_fit = None
    if len(isc_list) >= 2:
        L_arr = np.array([p[0] for p in isc_list], dtype=float)
        I_arr = np.array([p[1] for p in isc_list], dtype=float)
        res = analyse_lsm(L_arr, I_arr, 'L', 'I_{sc}', 'lx', 'mA')
        isc_fit = {"L": L_arr, "I": I_arr, "res": res}
        lines.append("Isc-L 线性拟合：Isc = k·L，k = {} mA/lx，截距 = {} mA，r = {}".format(
            formatted(res.m, 4), formatted(res.b), formatted(res.r, 4)))
        lines.append("短路电流与光照强度成正比（Isc ∝ L），符合光伏效应规律。")
    else:
        lines.append("提示：Isc（Ur/R₀）有效点不足 2 个，未做 Isc-L 线性拟合。")
    if len(uoc_list) >= 2:
        L_arr = np.array([p[0] for p in uoc_list], dtype=float)
        U_arr = np.array([p[1] for p in uoc_list], dtype=float)
        A, B = _log_fit(L_arr, U_arr)
        if A is not None:
            uoc_fit = {"L": L_arr, "U": U_arr, "A": A, "B": B}
            lines.append("Uoc-L 对数拟合：{} (V)".format(_log_eq("Uoc", A, B)))
    else:
        lines.append("提示：Uoc 有效点不足 2 个，未做 Uoc-L 对数拟合。")
    return {"lines": lines, "records": records,
            "isc_fit": isc_fit, "uoc_fit": uoc_fit}


def _calc_u_l(rows, fit_mode, var):
    """表 4/表 5 输出电压与光照：对数拟合（有载）或线性拟合（反偏）。"""
    d_arr, U_arr = _xy_pairs(rows, "c0", "c1")
    if len(d_arr) < 2:
        raise ValueError("至少需要 2 组有效数据（d 与 U 均填写）。")
    if (d_arr <= 0).any():
        raise ValueError("灯距 d 需要为大于 0 的数值。")
    L_arr = np.array([_light_lx(v) for v in d_arr], dtype=float)
    if fit_mode == "linear":
        k, b = np.polyfit(L_arr, U_arr, 1)
        r = float(np.corrcoef(L_arr, U_arr)[0, 1])
        lines = ["线性拟合：{}".format(_linear_eq(var, float(k), b))]
        lines.append("反向偏压下光电流 Iph ∝ L，输出电压与光照近似成正比。")
        return {"lines": lines, "d": d_arr, "L": L_arr, "U": U_arr,
                "fit_mode": fit_mode, "k": float(k), "b": float(b), "r": r}
    A, B = _log_fit(L_arr, U_arr)
    if A is None:
        raise ValueError("对数拟合失败（需要至少 2 个有效点）。")
    lines = ["对数拟合：{}".format(_log_eq(var, A, B))]
    lines.append("接负载时输出电压随光照强度近似按对数规律增长。")
    return {"lines": lines, "d": d_arr, "L": L_arr, "U": U_arr,
            "fit_mode": fit_mode, "A": A, "B": B}


# ── schema / preview / handle_structured ────────────────────────────────


def schema():
    dark = load_sample_data_numeric("exp15", "硅光电池暗伏安特性测量（示例数据）")
    output = load_sample_data_numeric("exp15", "硅光电池输出特性测量（示例数据）")
    uoc_isc = load_sample_data_numeric("exp15", "硅光电池开路电压、短路电流与光照特性测量（示例数据）")

    table4_sample = (
        ["50", "0.330"], ["45", "0.336"], ["40", "0.343"], ["35", "0.350"],
        ["30", "0.357"], ["25", "0.366"], ["20", "0.377"],
    )
    table5_sample = (
        ["50", "4.00"], ["45", "4.94"], ["40", "6.25"], ["35", "8.16"],
        ["30", "11.11"], ["25", "16.00"], ["20", "25.00"],
    )

    table1 = make_table(
        "table1", "硅光电池暗伏安特性测量数据表",
        ["U(V)", "I(mA)"],
        sample=dark,
        initial_rows=6,
        description="无光照、正向偏压下测量工作电流 I 与工作电压 U（正向偏压一般在 0.1~0.5 V 量级）。",
        chart={"x_column": "c0", "y_column": "c1", "x_label": "正向电压 U (V)",
               "y_label": "工作电流 I (mA)", "title": "暗伏安特性 I-U 曲线",
               "fit": "exponential"},
    )
    table1["calc"] = {"label": "拟合暗伏安特性"}

    table2 = make_table(
        "table2", "硅光电池输出特性测量数据表",
        ["R(Ω)", "U_20cm(V)", "U_30cm(V)", "U_40cm(V)", "U_50cm(V)"],
        sample=output,
        initial_rows=11,
        description="无偏压、以电阻箱作负载（负载电阻不小于 50 Ω，电压表用电压档），"
                    "分别测量 d=20/30/40/50cm 时负载工作电压 U。",
        chart={"x_column": "c0", "y_column": "c4", "x_label": "负载电阻 R (Ω)",
               "y_label": "工作电压 U (V)", "title": "输出特性 U-R 曲线（d=50cm）",
               "fit": "auto"},
    )
    table2["calc"] = {"label": "计算输出特性参数"}

    table3 = make_table(
        "table3", "开路电压、短路电流与光照特性测量数据表",
        ["d(cm)", "Uoc(V)", "Ur(mV)", "Isc(mA)"],
        sample=uoc_isc,
        readonly=(3,),
        initial_rows=7,
        description="不同灯距 d 下测量开路电压 Uoc 与取样电阻电压 Ur，短路电流 Isc=Ur/R₀ 自动换算。",
    )
    table3["calc"] = {"label": "拟合 Isc/Uoc-光照特性"}

    table4 = make_table(
        "table4", "输出电压与光照特性测量数据表（进阶）",
        ["d(cm)", "U(V)"],
        sample=table4_sample,
        required=False,
        initial_rows=5,
        description="接负载 R_L，测量不同灯距 d 下负载工作电压 U。",
    )
    table4["calc"] = {"label": "拟合 U-L 特性"}

    table5 = make_table(
        "table5", "反向偏压下输出电压与光照特性测量数据表（高阶）",
        ["d(cm)", "U(V)"],
        sample=table5_sample,
        required=False,
        initial_rows=5,
        description="反向偏压下，测量不同灯距 d 的负载工作电压 U。",
    )
    table5["calc"] = {"label": "拟合反向偏压 U-L 特性"}

    return make_schema(
        (
            "本实验测量硅光电池的光电特性。请按指导书完成以下数据处理："
            "① 暗伏安特性（无光照、正向偏压下的 I–U 关系）；② 输出特性（无偏压、电阻箱作负载）；"
            "③ 不同灯距下的开路电压 " r"$U_{oc}$" " 与短路电流（" r"$I_{sc} = U_r/R_0$" "）"
            "及光照特性；④ 接负载时的输出电压-光照特性；⑤ 反向偏压下负载工作电压与灯距的关系。"
        ),
        [
            table1, table2, table3, table4, table5,
        ],
        parameters=[
            {"id": "R0", "label": "取样电阻 R₀（换算短路电流）", "unit": "Ω",
             "type": "number", "default": 1.0, "required": True,
             "help": "Isc = Ur/R₀，按实际取样电阻标注值填写（常为 1 Ω）。"},
            {"id": "RL", "label": "负载电阻 R_L（输出电压测量）", "unit": "Ω",
             "type": "number", "default": 1000.0, "required": False,
             "help": "进阶/高阶内容所用负载电阻箱阻值（不小于 50 Ω）。"},
        ],
        analysis_hints="暗伏安特性应呈指数增长，拟合 b=q/kT 约 30~40 1/V（对应结温约 300 K）；"
                       "输出特性 P=U²/R 存在最大功率点 Pm（对应最佳负载 Rm），负载电阻不小于 50 Ω；"
                       "Isc 与光照 L 近似成正比（Isc-L 线性），Uoc 与 ln L 近似成正比；"
                       "L = 40·(50/d)² lx；反向偏压下输出电压与光照近似成正比；"
                       "填充因子 FF = Pm/(Uoc·Isc) 越大光电转换品质越好。",
        preview_enabled=True,
        report_enabled=False,
        formulas=get_formulas("exp15"),
        variables=get_variables("exp15"),
        table_theory=get_table_theory("exp15"),
    )


def _compute_all(payload):
    """按 payload 计算各表结果并补全表 3 的短路电流列 Isc = Ur/R₀。

    返回富计算结果（含 numpy 数组与最小二乘对象，供 handle_structured 使用），
    其中 calc_results[tid]["lines"] 为可直接展示的文本结论。
    """
    params = payload.get("parameters") or {}
    r0 = as_number(params.get("R0"))
    r0_ok = r0 is not None and r0 > 0
    tables = copied_tables(payload)
    calc_results = {}
    calc_messages = []
    fit_notes = {}

    # 表 3 只读列：Isc = Ur/R₀
    for row in tables.get("table3", []):
        ur = as_number(row.get("c2"))
        row["c3"] = formatted(ur / r0, 4) if r0_ok and ur is not None else ""

    def attempt(table_id, filled_cols, fn, notes):
        rows = tables.get(table_id) or []
        if not any(_has_value(r.get(c)) for r in rows for c in filled_cols):
            return
        try:
            calc_results[table_id] = fn()
            fit_notes[table_id] = notes
        except ValueError as exc:
            calc_messages.append(str(exc))

    attempt("table1", ("c0", "c1"),
            lambda: _calc_dark(tables["table1"]),
            ["暗状态下等效二极管：I = I₀·(e^{qU/kT} − 1)。",
             "线性化：ln I = ln I₀ + (q/kT)·U，由斜率 b 得热电压 kT/q 与结温 T。"])

    attempt("table2", ("c0", "c1", "c2", "c3", "c4"),
            lambda: _calc_output(tables["table2"]),
            ["输出功率 P = U·I = U²/R，最大功率点对应最佳负载 Rm。",
             "填充因子 FF = Pm/(Uoc·Isc)，反映电池光电转换品质。",
             "负载电阻不小于 50 Ω，电压表需用电压档。"])

    attempt("table3", ("c0", "c1", "c2"),
            lambda: _calc_uoc_isc(tables["table3"], r0 if r0_ok else None),
            ["短路电流 Isc = Ur/R₀，与光照 L 近似成正比。",
             "开路电压 Uoc = (kT/q)·ln(1+Isc/I₀)，与 ln L 近似成正比。",
             "光照强度换算：L = 40·(50/d)² lx。"])

    attempt("table4", ("c0", "c1"),
            lambda: _calc_u_l(tables["table4"], "log", "U"),
            ["接负载 R_L 时输出电压 U 与 ln L 近似成正比。"])

    attempt("table5", ("c0", "c1"),
            lambda: _calc_u_l(tables["table5"], "linear", "U"),
            ["反向偏压下光电流 Iph ∝ L，输出电压 U = Iph·R_L 与 L 近似成正比。"])

    return {"tables": tables, "calc_results": calc_results,
            "calc_messages": calc_messages, "fit_notes": fit_notes}


def preview(payload):
    """实时预计算接口：只返回前端需要且可 JSON 序列化的结论行。"""
    raw = _compute_all(payload)
    return {"tables": raw["tables"],
            "calc_results": {tid: {"lines": v["lines"]}
                             for tid, v in raw["calc_results"].items()},
            "calc_messages": raw["calc_messages"],
            "fit_notes": raw["fit_notes"]}


def handle_structured(workpath, payload):
    params = payload.get("parameters") or {}
    r0 = as_number(params.get("R0"))
    r0_ok = r0 is not None and r0 > 0
    rl = as_number(params.get("RL")) or 1000.0

    raw = _compute_all(payload)
    calc = raw["calc_results"]
    if not calc:
        return {"code": 1, "message": "；".join(raw["calc_messages"]) or "请先填写至少一张测量表。"}

    os.makedirs(workpath, exist_ok=True)
    warnings = list(raw["calc_messages"])
    summary = []
    charts = []

    docu = Document()
    style_doc_font(docu)
    docu.add_heading(name(), level=0)
    docu.add_paragraph("实验参数：取样电阻 R₀ = {} Ω；负载电阻 R_L = {} Ω".format(
        formatted(r0) if r0_ok else "未填写", rl))

    # ── 一、暗伏安特性 ────────────────────────────────────────────────
    dark = calc.get("table1")
    if dark:
        docu.add_heading("一、暗伏安特性", level=1)
        _add_doc_table(docu, ["U(V)", "I(mA)"],
                       [[formatted(u), formatted(i)] for u, i in zip(dark["U"], dark["I"])])
        x_fit = np.linspace(min(dark["U"]), max(dark["U"]), 200)
        y_fit = dark["i0"] * (np.exp(dark["b"] * x_fit) - 1)
        fn = _save_chart(workpath, "chart_dark_iv.png",
                         "无光照正向偏压时硅光电池的 I-U 特性曲线",
                         "U / V", "I / mA", dark["U"], dark["I"], x_fit, y_fit,
                         "指数拟合")
        charts.append({"filename": fn, "title": "暗伏安特性 I-U 曲线"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
        docu.add_paragraph("拟合方程：I = I₀·(e^{bU} − 1)")
        docu.add_paragraph("反向饱和电流 I₀ = {} mA".format(formatted(dark["i0"])))
        docu.add_paragraph("拟合系数 b = q/kT = {} 1/V".format(formatted(dark["b"])))
        docu.add_paragraph("热电压 kT/q = {} mV".format(formatted(dark["kt_q"] * 1000)))
        docu.add_paragraph("结温估计 T ≈ {} K".format(round(dark["temp"])))
        summary.extend(dark["lines"])
    else:
        warnings.append("未填写暗伏安特性测量表。")
    docu.add_paragraph()

    # ── 二、输出特性（全部光照距离列 + 与表 3 实测对照）──────────────
    out = calc.get("table2")
    if out:
        docu.add_heading("二、输出特性", level=1)
        series = out["series"]
        primary = next((s for s in series if s["label"] == "d=50cm"), series[0])
        _add_doc_table(docu, ["R(Ω)", "U(V)", "I(mA)", "P(mW)"],
                       [[formatted(r), formatted(u), formatted(i * 1000), formatted(p * 1000)]
                        for r, u, i, p in zip(primary["R"], primary["U"],
                                              primary["I"], primary["P"])])
        fn = _output_pr_chart(workpath, series, primary["label"])
        charts.append({"filename": fn, "title": "输出特性 P-R 曲线（全距离叠加）"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
        docu.add_paragraph("各光照距离下输出特性参数（以 {} 列为代表计算）：".format(primary["label"]))
        for s in series:
            docu.add_paragraph(
                "{}（L ≈ {:.1f} lx）：Uoc ≈ {} V，Isc ≈ {} mA，Pm = {} mW"
                "（Rm = {} Ω，Um = {} V，Im = {} mA），FF ≈ {}".format(
                    s["label"], s["L"], formatted(s["uoc"]), formatted(s["isc"] * 1000),
                    formatted(s["pm"] * 1000), formatted(s["rm"]),
                    formatted(s["um"]), formatted(s["im"] * 1000),
                    formatted(s["ff"]) if s["ff"] is not None else "—"))
        t3c = calc.get("table3")
        if t3c:
            match = next((x for x in t3c["records"]
                          if abs(x["d"] - 50.0) < 1e-9 and x["uoc"] is not None), None)
            if match:
                dev = (primary["uoc"] - match["uoc"]) / match["uoc"] * 100
                docu.add_paragraph(
                    "与表 3 实测对照（d=50cm）：本表近似 Uoc = {} V，实测 Uoc = {} V，"
                    "相对偏差 {:.1f}%".format(formatted(primary["uoc"]),
                                             formatted(match["uoc"]), dev))
                if abs(dev) > 10:
                    warnings.append("输出特性表近似 Uoc 与表 3 实测值偏差 {:.1f}%（>10%），"
                                    "请检查 R 最大档数据。".format(dev))
        summary.extend(out["lines"])
    else:
        warnings.append("未填写输出特性测量表。")
    docu.add_paragraph()

    # ── 三、开路电压、短路电流与光照 ──────────────────────────────────
    t3c = calc.get("table3")
    if t3c:
        docu.add_heading("三、开路电压、短路电流与光照特性", level=1)
        records = t3c["records"]
        _add_doc_table(docu, ["d(cm)", "L(lx)", "Uoc(V)", "Ur(mV)", "Isc(mA)"],
                       [[formatted(x["d"]), formatted(x["L"]), formatted(x["uoc"]),
                         formatted(x["ur"]), formatted(x["isc"])] for x in records])
        if t3c["isc_fit"]:
            f = t3c["isc_fit"]
            x_fit = np.linspace(f["L"].min(), f["L"].max(), 200)
            fn = _save_chart(workpath, "chart_isc_l.png",
                             "短路电流-光照特性曲线 Isc-L",
                             "L / lx", "Isc / mA", f["L"], f["I"],
                             x_fit, f["res"].m * x_fit + f["res"].b, "线性拟合")
            charts.append({"filename": fn, "title": "Isc-L 特性曲线"})
            docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
            docu.add_paragraph("短路电流与光照近似成正比：Isc = k·L")
            insert_data_lsm(docu, f["res"], "word")
        if t3c["uoc_fit"]:
            f = t3c["uoc_fit"]
            x_fit = np.linspace(f["L"].min(), f["L"].max(), 200)
            fn = _save_chart(workpath, "chart_uoc_l.png",
                             "开路电压-光照特性曲线 Uoc-L",
                             "L / lx", "Uoc / V", f["L"], f["U"],
                             x_fit, f["A"] * np.log(x_fit) + f["B"],
                             "对数拟合 Uoc=A·lnL+B")
            charts.append({"filename": fn, "title": "Uoc-L 特性曲线"})
            docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
            docu.add_paragraph("开路电压与光照的对数近似成正比：Uoc = A·ln L + B")
            docu.add_paragraph("A = {} V，B = {} V".format(formatted(f["A"]), formatted(f["B"])))
        summary.extend(t3c["lines"])
    else:
        warnings.append("未填写开路电压、短路电流与光照特性测量表。")
    docu.add_paragraph()

    # ── 四、五、输出电压与光照（进阶/高阶，选做表不填不警告）──────────
    for tid, heading, fname, chart_title in (
        ("table4", "四、输出电压与光照特性（进阶）", "chart_u_l.png",
         "输出电压-光照特性曲线 U-L"),
        ("table5", "五、反向偏压下输出电压与光照特性（高阶）", "chart_u_rev_l.png",
         "反向偏压下输出电压-光照特性曲线 U-L"),
    ):
        info = calc.get(tid)
        if not info:
            continue
        docu.add_heading(heading, level=1)
        L_arr, U_arr = info["L"], info["U"]
        _add_doc_table(docu, ["d(cm)", "L(lx)", "U(V)"],
                       [[formatted(d), formatted(L), formatted(u)]
                        for d, L, u in zip(info["d"], L_arr, U_arr)])
        x_fit = np.linspace(L_arr.min(), L_arr.max(), 200)
        if info["fit_mode"] == "linear":
            fit_y = info["k"] * x_fit + info["b"]
            fit_label = "线性拟合 U=k·L+b"
            fn = _save_chart(workpath, fname, chart_title, "L / lx", "U / V",
                             L_arr, U_arr, x_fit, fit_y, fit_label)
            docu.add_paragraph("线性拟合：U = k·L + b，k = {} V/lx，b = {} V".format(
                formatted(info["k"], 4), formatted(info["b"])))
        else:
            fit_y = info["A"] * np.log(x_fit) + info["B"]
            fit_label = "对数拟合 U=A·lnL+B"
            fn = _save_chart(workpath, fname, chart_title, "L / lx", "U / V",
                             L_arr, U_arr, x_fit, fit_y, fit_label)
            docu.add_paragraph("对数拟合：U = A·ln L + B，A = {} V，B = {} V".format(
                formatted(info["A"]), formatted(info["B"])))
        charts.append({"filename": fn, "title": chart_title})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
        summary.extend(info["lines"])
        docu.add_paragraph()

    docu.save(os.path.join(workpath, name() + ".docx"))

    return {
        "code": 0,
        "summary": summary or ["硅光电池光电特性数据已处理，详见 Word 文档。"],
        "warnings": warnings,
        "charts": charts,
        "document": name() + ".docx",
    }


def handle(workpath, extension):
    """旧版单表入口（兼容 b_adapter 等旧调用路径）：暗伏安特性 I-U 拟合。"""
    try:
        excelpath = workpath + name() + '.' + extension
        if extension == 'csv':
            with open(excelpath, 'rb') as f:
                encode = chardet.detect(f.read())["encoding"]
            data = pd.read_csv(excelpath, header=0, names=["U", "I"], encoding=encode)
        else:
            data = pd.read_excel(excelpath, header=0, names=["U", "I"])

        os.remove(excelpath)

        U = np.asarray([float(v) for v in data['U']], dtype=float)
        I = np.asarray([float(v) for v in data['I']], dtype=float)
        i0, b = _fit_diode(U, I)
        x_fit = np.linspace(U.min(), U.max(), 200) if len(U) else np.array([])
        y_fit = i0 * (np.exp(b * x_fit) - 1) if i0 is not None else np.array([])

        imgpath = workpath + "1.jpg"
        _save_chart(workpath, "1.jpg",
                    "无光照正向偏压时硅光电池的 I-U 特性曲线",
                    "U / V", "I / mA", U, I, x_fit, y_fit,
                    "指数拟合")

        docu = Document()
        style_doc_font(docu)
        docu.add_paragraph(name())
        docu.add_paragraph('无光照正向偏压时硅光电池的 I-U 特性曲线')
        docu.add_picture(imgpath)
        if i0 is not None and b is not None:
            docu.add_paragraph('反向饱和电流 I₀ = {} mA'.format(formatted(i0)))
            docu.add_paragraph('拟合系数 b = {} 1/V'.format(formatted(b)))
        docu.save(workpath + name() + ".docx")
        os.remove(imgpath)
        return 0
    except Exception:
        traceback.print_exc()
        return 1
