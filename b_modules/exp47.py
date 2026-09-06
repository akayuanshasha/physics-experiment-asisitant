"""接触角仪实验数据处理模块。

实验内容（指导书，论文格式报告）：
- 基础：影像分析法测量纯水在不同固体材料（自备树叶 + 另一种平整材料）
  表面的静态接触角，每种材料测 3 次，判断浸润特性；
- 提升：SiO2 微纳涂料 / 蜡烛烟灰 / 低表面能涂料三种方法制作疏水玻璃，
  与普通玻璃对比；
- 进阶：液滴体积增、减法或倾斜板法测量动态接触角（前进角 θA、后退角 θR、
  起始滚动角 α），计算接触角滞后；
- 高阶：用接触角评估固体表面洁净度。

网页流程（新版结构化接口）：
1. schema() 声明静态接触角、θ/2 法几何校核、动态接触角三张表与实验参数；
2. preview() 实时给出每种材料的平均接触角、标准差、不确定度与浸润性分类；
3. handle_structured() 生成分材料接触角柱状图、几何校核图、滞后图与 Word 报告。

旧版 CSV 流程 handle() 保留，供旧接口与 AI 助教调用。
"""

from __future__ import annotations

import math
import os
import traceback
from typing import Any

import numpy as np

from structured_support import as_number, copied_tables, make_schema, make_table, structured_result
from theory_content import get_table_theory

# ── 常数与判据 ──
_GAMMA_WATER = 72.8        # 20 °C 纯水表面张力（mN/m）
_DELTA_INST_DEFAULT = 0.5  # 接触角仪器允差默认值（°）
_CONFIDENCE_P = 0.95       # 不确定度置信概率
_SYM_LIMIT = 5.0           # 左右接触角不对称提示阈值（°）
# t 分布系数（P=0.95，自由度 ν = n-1），用于 A 类不确定度；超出范围取 2.0
_T95 = {1: 12.71, 2: 4.30, 3: 3.18, 4: 2.78, 5: 2.57, 6: 2.45, 7: 2.36,
        8: 2.31, 9: 2.26, 10: 2.23, 11: 2.20, 12: 2.18, 13: 2.16, 14: 2.14,
        15: 2.13, 16: 2.12, 17: 2.11, 18: 2.10, 19: 2.09, 20: 2.09,
        25: 2.06, 30: 2.04}


def name():
    return "接触角仪"


# ──────────────────────────────────────────────
# 基础工具
# ──────────────────────────────────────────────

def _has_value(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _param_number(parameters: dict[str, Any], key: str, label: str,
                  default: float | None = None) -> float:
    """读取实验参数；未填写返回 default（无 default 时报错），非数字报错。"""
    value = parameters.get(key)
    if not _has_value(value):
        if default is not None:
            return default
        raise ValueError(f"{label}未填写")
    number = as_number(value)
    if number is None:
        raise ValueError(f"{label}必须是数字")
    return number


def _t_factor(count: int) -> float:
    return _T95.get(count - 1, 2.0)


def _wetting_class(theta: float) -> str:
    if theta < 5:
        return "超亲水（θ<5°）"
    if theta < 90:
        return "亲水（θ<90°）"
    if theta <= 150:
        return "疏水（90°<θ≤150°）"
    return "超疏水（θ>150°）"


def _mean_sigma_unc(values: list[float], delta_inst: float) -> tuple[float, float, float]:
    """平均值、样本标准差与扩展不确定度（P=0.95，仪器允差按均匀分布 C=√3）。"""
    count = len(values)
    mean = sum(values) / count
    sigma = 0.0
    if count > 1:
        sigma = math.sqrt(sum((value - mean) ** 2 for value in values) / (count - 1))
    u_a = _t_factor(count) * sigma / math.sqrt(count)
    u_b = delta_inst / math.sqrt(3)
    unc = math.sqrt(u_a ** 2 + u_b ** 2)
    return mean, sigma, unc


# ──────────────────────────────────────────────
# 各表计算
# ──────────────────────────────────────────────

def _calc_static(rows: list[dict[str, Any]], gamma: float, delta_inst: float) -> dict[str, Any]:
    """静态接触角：按材料分组统计，输出浸润性分类、粘附功与左右角对称性检查。"""
    groups: dict[str, list[tuple[float, float]]] = {}  # 材料 -> [(θ̄, |θL−θR|), ...]
    asym_rows: list[str] = []
    used = False
    for index, row in enumerate(rows or [], start=1):
        if not any(_has_value(row.get(column)) for column in ("c0", "c1", "c2", "c3")):
            continue
        material = str(row.get("c0") or "").strip()
        raw_left = row.get("c1")
        raw_right = row.get("c2")
        raw_mean = row.get("c3")
        if not material or not _has_value(raw_left) or not _has_value(raw_right):
            raise ValueError("静态接触角表第 {} 行需填写材料名称、左接触角和右接触角".format(index))
        try:
            left = float(raw_left)
            right = float(raw_right)
            mean = float(raw_mean) if _has_value(raw_mean) else (left + right) / 2
        except (TypeError, ValueError):
            raise ValueError("静态接触角表第 {} 行角度数据不是数字".format(index))
        if not (0 <= left <= 180 and 0 <= right <= 180 and 0 <= mean <= 180):
            raise ValueError("静态接触角表第 {} 行接触角应在 0~180° 之间".format(index))
        groups.setdefault(material, []).append((mean, abs(left - right)))
        if abs(left - right) > _SYM_LIMIT:
            asym_rows.append(str(index))
        used = True
    if not used:
        raise ValueError("请填写静态接触角测量表（材料、左接触角、右接触角）")

    lines: list[str] = []
    materials: list[dict[str, Any]] = []
    for material, entries in groups.items():
        means = [entry[0] for entry in entries]
        mean, sigma, unc = _mean_sigma_unc(means, delta_inst)
        wetting = _wetting_class(mean)
        adhesion_work = gamma * (1 + math.cos(math.radians(mean)))
        adhesion_tension = gamma * math.cos(math.radians(mean))
        lines.append(
            "材料「{}」共 {} 次：θ̄ = {:.1f}°，σ = {:.2f}°，U = {:.1f}°（P=0.95）→ {}".format(
                material, len(means), mean, sigma, unc, wetting))
        lines.append(
            "　　粘附功 W_a = γ_lv(1+cosθ̄) = {:.1f} mN/m；粘附张力 γ_lv·cosθ̄ = {:+.1f} mN/m".format(
                adhesion_work, adhesion_tension))
        materials.append({
            "name": material, "mean": mean, "unc": unc,
            "count": len(means), "wetting": wetting,
        })

    if asym_rows:
        lines.append(
            "注意：第 {} 行左右接触角相差超过 {}°，样品表面可能不平整、受污染或基线未对准。".format(
                "、".join(asym_rows), _SYM_LIMIT))
    if len(groups) > 1:
        ranked = sorted(materials, key=lambda item: item["mean"], reverse=True)
        lines.append("各材料接触角排序：" + " ＞ ".join(
            "{} {:.1f}°".format(item["name"], item["mean"]) for item in ranked))
        if any(("清洁" in name) or ("清洗" in name) or ("污染" in name) for name in groups):
            lines.append(
                "高阶内容提示：对玻璃、金属等高能表面，接触角越小通常表面越洁净，"
                "可据此比较清洁前后样品的洁净度。")
    return {"lines": lines, "materials": materials}


def _calc_geometry(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """θ/2 法几何校核：θ = 2·arctan(2h/d)，与仪器测得值对比。"""
    theta_inst: list[float] = []
    theta_calc: list[float] = []
    used = False
    for index, row in enumerate(rows or [], start=1):
        if not any(_has_value(row.get(column)) for column in ("c0", "c1", "c2")):
            continue
        try:
            measured = float(row.get("c0"))
            height = float(row.get("c1"))
            diameter = float(row.get("c2"))
        except (TypeError, ValueError):
            raise ValueError("几何校核表第 {} 行数据不完整或不是数字".format(index))
        if diameter <= 0:
            raise ValueError("几何校核表第 {} 行接触直径 d 必须大于 0".format(index))
        if not 0 <= measured <= 180:
            raise ValueError("几何校核表第 {} 行仪器接触角应在 0~180° 之间".format(index))
        theta_inst.append(measured)
        theta_calc.append(2 * math.degrees(math.atan(2 * height / diameter)))
        used = True
    if not used:
        raise ValueError("请填写几何校核表（仪器接触角、液滴高度 h、接触直径 d）")

    diffs = [measured - calc for measured, calc in zip(theta_inst, theta_calc)]
    mean_diff = sum(diffs) / len(diffs)
    mean_abs_diff = sum(abs(diff) for diff in diffs) / len(diffs)
    max_abs_diff = max(abs(diff) for diff in diffs)
    values_text = "、".join("{:.1f}°".format(value) for value in theta_calc) \
        if len(theta_calc) <= 12 else "共 {} 组".format(len(theta_calc))
    lines = [
        "θ/2 法算得各组接触角：{}".format(values_text),
        "与仪器值对比：平均偏差 {:+.2f}°，平均绝对偏差 {:.2f}°，最大绝对偏差 {:.2f}°".format(
            mean_diff, mean_abs_diff, max_abs_diff),
    ]
    if abs(mean_diff) <= 3:
        lines.append("两种方法结果一致，θ/2 法球冠近似成立，测量可靠。")
    else:
        lines.append(
            "两种方法偏差较大：检查 h、d 读数是否正确，液滴是否过大"
            "（建议 2~5 μl，球冠近似失效时 θ/2 法偏差增大）。")
    return {"lines": lines, "theta_inst": theta_inst, "theta_calc": theta_calc}


def _calc_dynamic(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """动态接触角：接触角滞后 Δθ = θ_A − θ_R 与起始滚动角。"""
    entries: list[dict[str, Any]] = []
    used = False
    for index, row in enumerate(rows or [], start=1):
        if not any(_has_value(row.get(column)) for column in ("c0", "c1", "c2", "c3")):
            continue
        material = str(row.get("c0") or "").strip() or "材料{}".format(index)
        try:
            advancing = float(row.get("c1"))
            receding = float(row.get("c2"))
        except (TypeError, ValueError):
            raise ValueError("动态接触角表第 {} 行前进角/后退角数据不是数字".format(index))
        rolling = None
        if _has_value(row.get("c3")):
            try:
                rolling = float(row.get("c3"))
            except (TypeError, ValueError):
                raise ValueError("动态接触角表第 {} 行滚动角数据不是数字".format(index))
        entries.append({
            "index": index, "name": material,
            "advancing": advancing, "receding": receding, "rolling": rolling,
        })
        used = True
    if not used:
        raise ValueError("请填写动态接触角表（材料、前进角 θ_A、后退角 θ_R）")

    lines: list[str] = []
    for entry in entries:
        hysteresis = entry["advancing"] - entry["receding"]
        suffix = ""
        if entry["rolling"] is not None:
            suffix = "；起始滚动角 α = {:.1f}°".format(entry["rolling"])
        if hysteresis < 0:
            suffix += "（注意：前进角小于后退角，请检查记录顺序）"
        lines.append(
            "材料「{}」：θ_A = {:.1f}°，θ_R = {:.1f}°，滞后 Δθ = {:.1f}°{}".format(
                entry["name"], entry["advancing"], entry["receding"], hysteresis, suffix))
    lines.append("接触角滞后越大，三相接触线越难移动，液滴越不易从表面滚落。")
    return {"lines": lines, "entries": entries}


# ──────────────────────────────────────────────
# 结构化接口
# ──────────────────────────────────────────────

def preview(payload: dict[str, Any]) -> dict[str, Any]:
    """网页输入时的实时预计算：静态接触角统计、θ/2 法校核、动态接触角滞后。"""
    tables = copied_tables(payload)
    parameters = payload.get("parameters") or {}
    calc_results: dict[str, Any] = {}
    calc_messages: list[str] = []
    fit_notes: dict[str, list[str]] = {}

    try:
        gamma = _param_number(parameters, "gamma_lv", "水的表面张力 γ_lv", _GAMMA_WATER)
        delta_inst = _param_number(parameters, "delta_inst", "接触角仪器允差 Δ_仪", _DELTA_INST_DEFAULT)
    except ValueError:
        gamma, delta_inst = _GAMMA_WATER, _DELTA_INST_DEFAULT

    static_rows = tables.get("static", []) or []
    if any(_has_value(row.get(column)) for row in static_rows for column in ("c0", "c1", "c2", "c3")):
        try:
            calc_results["static"] = _calc_static(static_rows, gamma, delta_inst)
            fit_notes["static"] = [
                "浸润性判据：θ<90° 亲水，θ>90° 疏水，θ>150° 超疏水。",
                "扩展不确定度 U = √((t·σ/√n)² + (Δ_仪/√3)²)，P=0.95。",
            ]
        except ValueError as exc:
            calc_messages.append(str(exc))

    geometry_rows = tables.get("geometry", []) or []
    if any(_has_value(row.get(column)) for row in geometry_rows for column in ("c0", "c1", "c2")):
        try:
            calc_results["geometry"] = _calc_geometry(geometry_rows)
            fit_notes["geometry"] = ["θ/2 法：小液滴近似球冠，θ = 2·arctan(2h/d)。"]
        except ValueError as exc:
            calc_messages.append(str(exc))

    dynamic_rows = tables.get("dynamic", []) or []
    if any(_has_value(row.get(column)) for row in dynamic_rows for column in ("c0", "c1", "c2", "c3")):
        try:
            calc_results["dynamic"] = _calc_dynamic(dynamic_rows)
            fit_notes["dynamic"] = ["接触角滞后 Δθ = θ_A − θ_R，滞后越大液滴越难滚落。"]
        except ValueError as exc:
            calc_messages.append(str(exc))

    return {"tables": tables, "calc_results": calc_results,
            "calc_messages": calc_messages, "fit_notes": fit_notes}


def _font_properties():
    from matplotlib.font_manager import FontProperties

    font_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "SourceHanSansSC-Regular.otf"))
    return FontProperties(fname=font_path) if os.path.isfile(font_path) else None


def _make_static_chart(workpath: str, info: dict[str, Any]) -> dict[str, Any]:
    """分材料静态接触角柱状图：误差棒 + 亲水/疏水/超疏水分区。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    font = _font_properties()
    materials = info["materials"]
    names = [item["name"] for item in materials]
    means = [item["mean"] for item in materials]
    uncs = [item["unc"] for item in materials]

    figure, axis = plt.subplots(figsize=(7.6, 5.0))
    for bottom, top, label, color, y_frac in (
            (0, 5, "超亲水", "#3a9d5d", 0.03),
            (5, 90, "亲水", "#2f7fc1", 0.26),
            (90, 150, "疏水", "#e6a23c", 0.67),
            (150, 180, "超疏水", "#d94b40", 0.92)):
        axis.axhspan(bottom, top, color=color, alpha=0.08, zorder=1)
        axis.text(0.995, y_frac, label, transform=axis.transAxes, ha="right",
                  va="center", fontsize=8, color="#666666", fontproperties=font)
    axis.axhline(90, color="#888888", linestyle="--", linewidth=0.8, zorder=2)
    axis.axhline(150, color="#888888", linestyle="--", linewidth=0.8, zorder=2)

    x = np.arange(len(names))
    axis.bar(x, means, yerr=uncs, capsize=5, color="#2f7fc1", alpha=0.85, zorder=3)
    for xi, item in enumerate(materials):
        axis.annotate("{:.1f}°±{:.1f}°".format(item["mean"], item["unc"]),
                      (xi, item["mean"] + item["unc"] + 4), ha="center",
                      fontsize=9, fontproperties=font)
    axis.set_xticks(x)
    if len(names) > 5:
        axis.set_xticklabels(names, rotation=20, ha="right", fontproperties=font, fontsize=9)
    else:
        axis.set_xticklabels(names, fontproperties=font, fontsize=9)
    axis.set_ylim(0, 180)
    axis.set_xlabel("材料", fontproperties=font)
    axis.set_ylabel("静态接触角 θ (°)", fontproperties=font)
    axis.set_title("纯水在不同材料表面的静态接触角", fontproperties=font)
    axis.grid(alpha=0.25, axis="y", zorder=0)
    figure.tight_layout()
    figure.savefig(os.path.join(workpath, "contact_angle_static.png"), dpi=220, bbox_inches="tight")
    plt.close(figure)
    return {"filename": "contact_angle_static.png", "title": "分材料静态接触角",
            "x_label": "材料", "y_label": "静态接触角 θ (°)"}


def _make_geometry_chart(workpath: str, info: dict[str, Any]) -> dict[str, Any]:
    """θ/2 法校核：θ_calc 对 θ_仪 散点 + 恒等线。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    font = _font_properties()
    theta_inst = info["theta_inst"]
    theta_calc = info["theta_calc"]
    lower = min(min(theta_inst), min(theta_calc)) - 2
    upper = max(max(theta_inst), max(theta_calc)) + 2

    figure, axis = plt.subplots(figsize=(7.6, 5.0))
    axis.scatter(theta_inst, theta_calc, s=40, color="#2f7fc1", zorder=4, label="测量点")
    axis.plot([lower, upper], [lower, upper], "--", color="#d94b40", linewidth=1.2,
              zorder=3, label="y = x（完全一致）")
    diffs = [measured - calc for measured, calc in zip(theta_inst, theta_calc)]
    mean_diff = sum(diffs) / len(diffs)
    axis.annotate(
        "平均偏差 θ̄_仪 − θ̄_calc = {:+.2f}°\n（球冠近似成立时两点应落在恒等线上）".format(mean_diff),
        xy=(0.04, 0.96), xycoords="axes fraction", va="top", fontsize=10,
        fontproperties=font,
        bbox=dict(boxstyle="round,pad=0.4", fc="#fff7e6", ec="#e6c88a", alpha=0.9),
    )
    axis.set_xlim(lower, upper)
    axis.set_ylim(lower, upper)
    axis.set_xlabel("仪器测得 θ_仪 (°)", fontproperties=font)
    axis.set_ylabel("θ/2 法算得 θ_calc (°)", fontproperties=font)
    axis.set_title("θ/2 法几何校核", fontproperties=font)
    axis.grid(alpha=0.25)
    if font is not None:
        axis.legend(prop=font)
    else:
        axis.legend()
    figure.tight_layout()
    figure.savefig(os.path.join(workpath, "contact_angle_geometry.png"), dpi=220, bbox_inches="tight")
    plt.close(figure)
    return {"filename": "contact_angle_geometry.png", "title": "θ/2 法几何校核",
            "x_label": "仪器接触角 θ_仪 (°)", "y_label": "θ/2 法接触角 θ_calc (°)"}


def _make_dynamic_chart(workpath: str, info: dict[str, Any]) -> dict[str, Any]:
    """动态接触角：前进角/后退角分组柱状图 + 滞后标注。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    font = _font_properties()
    entries = info["entries"]
    names = [entry["name"] for entry in entries]
    advancing = [entry["advancing"] for entry in entries]
    receding = [entry["receding"] for entry in entries]

    figure, axis = plt.subplots(figsize=(7.6, 5.0))
    x = np.arange(len(names))
    width = 0.34
    axis.bar(x - width / 2, advancing, width, color="#d94b40", alpha=0.9, zorder=3, label="前进角 θ_A")
    axis.bar(x + width / 2, receding, width, color="#2f7fc1", alpha=0.9, zorder=3, label="后退角 θ_R")
    for xi, entry in enumerate(entries):
        hysteresis = entry["advancing"] - entry["receding"]
        axis.annotate("Δθ = {:.1f}°".format(hysteresis),
                      (xi, max(entry["advancing"], entry["receding"]) + 3),
                      ha="center", fontsize=9, color="#3a9d5d", fontproperties=font)
    axis.set_xticks(x)
    axis.set_xticklabels(names, fontproperties=font, fontsize=9)
    axis.set_ylabel("接触角 (°)", fontproperties=font)
    axis.set_xlabel("材料", fontproperties=font)
    axis.set_title("动态接触角与滞后", fontproperties=font)
    axis.set_ylim(0, max(max(advancing), max(receding)) * 1.2)
    axis.grid(alpha=0.25, axis="y", zorder=0)
    if font is not None:
        axis.legend(prop=font)
    else:
        axis.legend()
    figure.tight_layout()
    figure.savefig(os.path.join(workpath, "contact_angle_dynamic.png"), dpi=220, bbox_inches="tight")
    plt.close(figure)
    return {"filename": "contact_angle_dynamic.png", "title": "动态接触角与滞后",
            "x_label": "材料", "y_label": "接触角 (°)"}


def handle_structured(workpath: str, payload: dict[str, Any]) -> dict[str, Any]:
    """新版结构化接口：统计静态接触角、校核 θ/2 法、计算滞后，输出摘要与图表。"""
    parameters = payload.get("parameters") or {}
    try:
        _param_number(parameters, "gamma_lv", "水的表面张力 γ_lv")
        _param_number(parameters, "delta_inst", "接触角仪器允差 Δ_仪")
        if _has_value(parameters.get("temperature")):
            _param_number(parameters, "temperature", "实验温度")
    except ValueError as exc:
        return {"code": 1, "message": str(exc)}

    result = preview(payload)
    static_info = result["calc_results"].get("static")
    if static_info is None:
        return {"code": 1, "message": "；".join(result["calc_messages"]) or "请先填写静态接触角测量表。"}

    os.makedirs(workpath, exist_ok=True)
    charts = [_make_static_chart(workpath, static_info)]
    geometry_info = result["calc_results"].get("geometry")
    if geometry_info is not None:
        charts.append(_make_geometry_chart(workpath, geometry_info))
    dynamic_info = result["calc_results"].get("dynamic")
    if dynamic_info is not None:
        charts.append(_make_dynamic_chart(workpath, dynamic_info))

    summary: list[str] = []
    for table_id, label in (("static", "静态接触角统计"),
                            ("geometry", "θ/2 法几何校核"),
                            ("dynamic", "动态接触角与滞后")):
        if table_id in result["calc_results"]:
            summary.append("【{}】".format(label))
            summary.extend(result["calc_results"][table_id]["lines"])

    enriched = dict(payload)
    enriched["tables"] = result["tables"]
    return structured_result(
        workpath, name(), schema(), enriched,
        summary=summary,
        charts=charts,
    )


# ──────────────────────────────────────────────
# 表结构与示例数据
# ──────────────────────────────────────────────

def _set_units(table: dict[str, Any], units: tuple[str, ...]) -> dict[str, Any]:
    """给 make_table 生成的列补充单位。"""
    for column, unit in zip(table["columns"], units):
        if unit:
            column["unit"] = unit
    return table


def schema() -> dict[str, Any]:
    static = _set_units(
        make_table(
            "static", "静态接触角测量表（基础/提升/高阶内容）",
            ["材料名称", "左接触角 θ_L", "右接触角 θ_R", "仪器平均 θ̄_仪（可留空）"],
            sample=[
                ["普通玻璃", 20.5, 22.0, 21.3],
                ["普通玻璃", 21.0, 21.8, 21.4],
                ["普通玻璃", 20.2, 21.5, 20.9],
                ["荷叶", 145.2, 146.8, 146.0],
                ["荷叶", 146.5, 147.9, 147.2],
                ["荷叶", 144.8, 146.2, 145.5],
            ],
            text_columns=(0,), min_rows=3, initial_rows=6,
            description=(
                "每种材料按指导书要求测量 3 次，记录仪器拟合的左、右接触角与平均值。"
                "材料如：自备树叶、荷叶、雨伞布、普通玻璃；提升内容：SiO2 涂料、蜡烛烟灰、"
                "低表面能涂料制作的疏水玻璃；高阶内容：清洁前后的玻璃样品。"
                "θ̄_仪 留空时自动取左右角平均。"
            ),
        ),
        ("", "°", "°", "°"),
    )
    static["calc"] = {"label": "统计接触角与浸润性"}

    geometry = _set_units(
        make_table(
            "geometry", "θ/2 法几何校核表",
            ["仪器接触角 θ_仪", "液滴高度 h", "接触直径 d"],
            sample=[
                [45.0, 0.621, 3.0],
                [46.0, 0.637, 3.0],
                [44.5, 0.614, 3.0],
                [45.5, 0.629, 3.0],
                [45.0, 0.621, 3.0],
            ],
            min_rows=1, initial_rows=5, required=False,
            description=(
                "从测量截图读取液滴高度 h 与固-液接触直径 d（停滴法 θ/2 的几何量），"
                "与仪器拟合接触角互相校核。θ=45° 的球冠液滴约 h=0.62 mm、d=3.0 mm。"
            ),
            chart={
                "x_column": "c1", "y_column": "c0",
                "x_label": "液滴高度 h (mm)", "y_label": "仪器接触角 θ_仪 (°)",
                "title": "接触角与液滴高度关系", "fit": "auto",
            },
        ),
        ("°", "mm", "mm"),
    )
    geometry["calc"] = {"label": "几何校核"}

    dynamic = _set_units(
        make_table(
            "dynamic", "动态接触角测量表（进阶内容，选做）",
            ["材料名称", "前进接触角 θ_A", "后退接触角 θ_R", "起始滚动角 α（可留空）"],
            sample=[
                ["荷叶", 151.0, 145.0, 8.5],
                ["雨伞布", 148.0, 140.5, ""],
            ],
            text_columns=(0,), min_rows=1, initial_rows=2, required=False,
            description=(
                "液滴体积增、减法或倾斜板法测量前进接触角 θ_A、后退接触角 θ_R；"
                "倾斜板法可另记录液滴开始滑动（滚动）时的倾斜角 α。"
            ),
        ),
        ("", "°", "°", "°"),
    )
    dynamic["calc"] = {"label": "计算接触角滞后"}
    dynamic["enabled_by_default"] = False

    parameters = [
        {"id": "gamma_lv", "label": "水的表面张力 γ_lv", "unit": "mN/m", "type": "number",
         "default": _GAMMA_WATER, "min": 0, "step": "any", "required": True,
         "help": "20 °C 纯水为 72.8 mN/m，可按实验温度查表修正"},
        {"id": "temperature", "label": "实验温度", "unit": "°C", "type": "number",
         "default": 20.0, "min": 0, "step": "any", "required": False,
         "help": "记录测试环境温度，用于说明 γ_lv 取值"},
        {"id": "delta_inst", "label": "接触角仪器允差 Δ_仪", "unit": "°", "type": "number",
         "default": _DELTA_INST_DEFAULT, "min": 0, "step": "any", "required": True,
         "help": "按仪器说明书或分辨力填写，B 类不确定度按均匀分布 C=√3"},
    ]

    theory = get_table_theory("exp47")
    return make_schema(
        (
            "本实验用影像分析法（停滴法 θ/2）测量纯水在固体表面的静态接触角，"
            "判断材料浸润特性（θ>90° 疏水、θ>150° 超疏水）；提升内容对比三种方法"
            "制作的疏水玻璃；进阶内容测量动态接触角与滞后；高阶内容用接触角评估表面洁净度。"
        ),
        [static, geometry, dynamic],
        parameters=parameters,
        parameters_sample={"gamma_lv": _GAMMA_WATER, "temperature": 20.0,
                           "delta_inst": _DELTA_INST_DEFAULT},
        analysis_hints=(
            "每种材料 3 次测量，检查数据离散程度（σ 偏大说明表面不均匀或操作不稳）；"
            "左右接触角相差大提示样品表面不平或基线未对准；"
            "θ/2 法仅对 2~5 μl 小液滴（球冠近似）成立；"
            "动态测量中前进角应大于后退角；"
            "高阶内容：对玻璃等高能表面，接触角越小通常洁净度越高。"
        ),
        preview_enabled=True,
        formulas=theory.get("_global", {}).get("formulas", []),
        variables=theory.get("_global", {}).get("variables", []),
        table_theory=theory,
    )


def _find_col(cols: list, exact_keys: tuple, contains_keys: tuple, fallback_index: int):
    """按精确匹配（忽略空格大小写）优先、包含匹配其次的顺序找列，找不到用回退列。

    原来的 `'h' in c` 会误匹配含字母 h 的 theta 表头，这里只做精确/带单位匹配。
    """
    normalized = [str(c).strip().lower().replace(" ", "") for c in cols]
    for key in exact_keys:
        if key in normalized:
            return cols[normalized.index(key)]
    for c, norm in zip(cols, normalized):
        if any(key in norm for key in contains_keys):
            return c
    return cols[fallback_index] if fallback_index < len(cols) else cols[0]


def handle(workpath: str, extension: str) -> int:
    """旧版 CSV 流程（θ、h、d 三列）：生成 Word 文档，供旧接口与 AI 助教调用。"""
    import chardet
    import pandas as pd
    from docx import Document
    from docx.oxml.ns import qn

    try:
        excelpath = workpath + name() + '.' + extension

        if extension == 'csv':
            with open(excelpath, 'rb') as f:
                encode = chardet.detect(f.read())["encoding"]
            data = pd.read_csv(excelpath, header=0, encoding=encode)
        else:
            data = pd.read_excel(excelpath, header=0)

        os.remove(excelpath)

        cols = list(data.columns)
        theta_col = _find_col(cols, ("theta", "θ", "接触角"), ("角度", "θ", "theta"), 0)
        h_col = _find_col(cols, ("h", "h(mm)", "hmm"), ("高度",), 1)
        d_col = _find_col(cols, ("d", "d(mm)", "dmm"), ("直径",), 2)

        theta = pd.to_numeric(data[theta_col], errors='coerce')
        h = pd.to_numeric(data[h_col], errors='coerce')
        d = pd.to_numeric(data[d_col], errors='coerce')
        valid = theta.notna() & h.notna() & d.notna()
        theta = theta[valid].reset_index(drop=True)
        h = h[valid].reset_index(drop=True)
        d = d[valid].reset_index(drop=True)
        if len(theta) == 0:
            raise ValueError("没有可用的数据行")

        theta_calc = 2 * np.degrees(np.arctan(2 * h / d))
        theta_mean, theta_sigma, _ = _mean_sigma_unc(list(theta), _DELTA_INST_DEFAULT)
        calc_mean, calc_sigma, _ = _mean_sigma_unc(list(theta_calc), _DELTA_INST_DEFAULT)
        diffs = [float(measured - calc) for measured, calc in zip(theta, theta_calc)]
        mean_diff = sum(diffs) / len(diffs)

        docu = Document()
        style_doc_font(docu)
        docu.add_heading(name(), level=0)
        docu.add_paragraph("接触角测量实验（旧版 CSV 流程）")
        docu.add_paragraph("水的表面张力 γ_lv = {:.1f} mN/m (20 °C)".format(_GAMMA_WATER))
        docu.add_paragraph()

        table = docu.add_table(rows=len(theta) + 1, cols=4, style='Table Grid')
        for j, header in enumerate(['θ_仪 (°)', 'h (mm)', 'd (mm)', 'θ_calc (°)']):
            table.rows[0].cells[j].text = header
        for i in range(len(theta)):
            table.rows[i + 1].cells[0].text = '{:.1f}'.format(theta.iloc[i])
            table.rows[i + 1].cells[1].text = '{:.3f}'.format(h.iloc[i])
            table.rows[i + 1].cells[2].text = '{:.3f}'.format(d.iloc[i])
            table.rows[i + 1].cells[3].text = '{:.1f}'.format(theta_calc.iloc[i])
        docu.add_paragraph()
        docu.add_paragraph("仪器接触角：θ̄ = {:.1f}°，σ = {:.2f}°".format(theta_mean, theta_sigma))
        docu.add_paragraph("θ/2 法接触角：θ̄_calc = {:.1f}°，σ = {:.2f}°".format(calc_mean, calc_sigma))
        docu.add_paragraph("平均偏差 θ̄_仪 − θ̄_calc = {:+.2f}°".format(mean_diff))
        docu.add_paragraph("浸润性判据（θ<90° 亲水，θ>90° 疏水，θ>150° 超疏水）：{}".format(
            _wetting_class(theta_mean)))
        docu.add_paragraph()
        docu.add_paragraph("公式：")
        docu.add_paragraph("Young方程：\\gamma_{sv} = \\gamma_{sl} + \\gamma_{lv}\\cos\\theta")
        docu.add_paragraph("润湿功：W_a = \\gamma_{lv}(1 + \\cos\\theta)")
        docu.add_paragraph("液滴近似：\\theta \\approx 2\\arctan(\\frac{2h}{d})")

        docu.save(workpath + name() + ".docx")
        return 0
    except Exception:
        traceback.print_exc()
        return 1
