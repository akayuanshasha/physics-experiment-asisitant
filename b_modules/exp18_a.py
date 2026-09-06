"""分光计的调节与使用实验模块
===================
一级大物光学实验 —— 分光计的调节和使用 B（光栅）

实验内容（对应《分光计的调节和使用B》实验指导）：
  a) 基础内容（仪器调节，本模块只处理数据部分）：
     1. 调节望远镜对平行光聚焦，且光轴垂直仪器主轴；
     2. 调节平行光管出射平行光，与望远镜共轴并垂直主轴；
     3. 调节光栅平面垂直平行光管光轴、刻痕与仪器转轴平行。
  b) 提升内容（测定光栅常数）：
     4. 以低压汞灯为光源，测绿光（546.07 nm）的 ±1 级衍射角
        φ+1、φ-1，由光栅方程 d·sinφ = kλ 求光栅常数 d。
  c) 进阶内容（测定未知光波长）：
     5. 同法测汞灯蓝紫线、两条黄线的 ±1 级衍射角，
        由 λ = d·sinφ 计算波长，并与标准值比较。
  d) 高阶内容（选做）：探究影响谱线分辨率的因素。
  实验指导说明"本实验不写实验报告"，故 report_enabled = False。

物理公式：
  光栅方程：        d·sinφ = kλ，k = 0, ±1, ±2, …
  衍射角：          φ = |θ_k − θ_0|（左右游标平均，消除偏心误差）
                    两侧对称检验：φ ≈ (φ+1 + φ-1)/2
  光栅常数：        d = λ_绿 / sin φ_绿   （k = 1，λ_绿 = 546.07 nm）
  波长：            λ = d·sinφ           （k = 1）
  不确定度约定（沿用原模块）：角度读数仪器允差 1′、估读允差 0.5′，
  C = √3、P = 0.95；d 与 λ 用 analyse_com 作误差传递。
"""

import re

from head import *  # 导入万能头
from docx.shared import Inches
from structured_support import (
    copied_tables, formatted, make_schema, make_table,
)
from theory_content import get_formulas, get_variables, get_table_theory
from sample_data_loader import load_sample_data
from optics_common import OpticsInputError, circular_signed_difference, parse_angle

# 汞灯谱线波长（nm）
_LAMBDA_GREEN = 546.07  # 绿线：指导书给定，用于求光栅常数
_LINE_THEORY = {  # 待测三条谱线的标准波长（汞灯），用于相对偏差比较
    "蓝紫": 435.83,
    "黄内": 576.96,
    "黄外": 579.07,
}
_LINE_ORDER = ("蓝紫", "黄内", "黄外")

# 分光计游标读数允差（沿用原模块约定）：仪器允差 1′、估读允差 0.5′
_DELTA_INSTRUMENT = 1.0 / 60.0
_DELTA_ESTIMATE = 0.5 / 60.0

# ±1 级衍射角一致性判定阈值（分）：两侧相差超过该值提示光栅未调好
_SYMMETRY_LIMIT_ARCMIN = 3.0
# 教学光栅常数合理范围（nm）：指导书提示常见 100~300 线/mm（d 约 3.3~10 μm），留余量
_D_MIN_NM = 1500.0
_D_MAX_NM = 15000.0
# 谱线波长与标准值的偏差警告阈值（%）
_DEV_WARN_PERCENT = 1.0

# 三次测量的 (θ1, θ2) 列对（θ1/θ2 为左、右游标）
_MEASUREMENT_COLS = (("c1", "c2"), ("c3", "c4"), ("c5", "c6"))
# 只读列：每次测量的衍射角与平均衍射角
_PHI_COLS = ("c7", "c8", "c9")
_PHI_MEAN_COL = "c10"

_TABLE_LABELS = [
    "谱线", "θ1(第1次)", "θ2(第1次)", "θ1(第2次)", "θ2(第2次)",
    "θ1(第3次)", "θ2(第3次)",
    "φ(第1次)/(°)", "φ(第2次)/(°)", "φ(第3次)/(°)", "平均衍射角 φ/(°)",
]

# 谱线行标签：如"绿(+1)"、"黄外(-1)"；"0级位置"行无级次
_LINE_ROW_PATTERN = re.compile(r"^(.+?)\s*[（(]\s*([+-])\s*1\s*[)）]\s*$")


def name():  # 返回实验名称
    return "分光计的调节与使用"


# ─────────────────────────────────────────────
# 数据解析辅助函数
# ─────────────────────────────────────────────
def _row_info(label):
    """解析谱线行标签。

    带级次的谱线行（绿(+1)、蓝紫(-1) 等）返回 ("line", 线名, "+"或"-")；
    "0级位置"行返回 ("zero", "", "")；其余返回 ("other", 原文, "")。
    """
    text = str(label or "").strip()
    match = _LINE_ROW_PATTERN.match(text)
    if match:
        return ("line", match.group(1).strip(), match.group(2))
    if "0" in text and "级" in text:
        return ("zero", "", "")
    return ("other", text, "")


def _find_zero_row(rows):
    """找到"0级位置"参考行，无则返回 None。"""
    for row in rows:
        if _row_info(row.get("c0", ""))[0] == "zero":
            return row
    return None


def _has_value(v):
    """判断单元格是否填写了内容。"""
    if v is None:
        return False
    if isinstance(v, str):
        return bool(v.strip())
    return True


def _reading(row, col):
    """解析单个游标读数（支持 "123°30′" 或十进制度），空/非法返回 None。"""
    raw = row.get(col, "")
    if raw is None or str(raw).strip() == "":
        return None
    try:
        return parse_angle(raw, "游标读数")
    except OpticsInputError:
        return None


def _phi_measurement(row, zero_row, m):
    """谱线行第 m 次测量的衍射角（度，取绝对值，左右游标平均）。

    φ = |θ_k − θ_0|（左、右游标分别求圆周差后取平均，消除刻度盘偏心误差），
    用 circular_signed_difference 处理 0°/360° 跨界。任一读数缺失返回 None。
    """
    if row is None or zero_row is None:
        return None
    t1_col, t2_col = _MEASUREMENT_COLS[m]
    diffs = []
    for col in (t1_col, t2_col):
        angle = _reading(row, col)
        zero = _reading(zero_row, col)
        if angle is None or zero is None:
            return None
        diffs.append(circular_signed_difference(zero, angle))
    return abs(sum(diffs) / len(diffs))


def _collect_phis_side(rows, keyword, side):
    """收集线名包含 keyword 的谱线某一级（"+"或"-"）的全部衍射角（度）。

    每行 3 次测量（每次已左右游标平均）。缺少 0 级位置行或该行数据不全时跳过。
    """
    zero_row = _find_zero_row(rows)
    if zero_row is None:
        return []
    values = []
    for row in rows:
        info = _row_info(row.get("c0", ""))
        if info[0] != "line" or keyword not in info[1] or info[2] != side:
            continue
        for m in range(3):
            phi = _phi_measurement(row, zero_row, m)
            if phi is not None:
                values.append(phi)
    return values


def _font():
    """返回思源黑体 FontProperties，找不到则返回 None（用默认字体）。"""
    path = os.environ.get('_B_FONT_PATH')
    if path and os.path.exists(path):
        return matplotlib.font_manager.FontProperties(fname=path)
    return None


def _decorate(ax, title, xlabel, ylabel):
    zhfont = _font()
    if zhfont:
        ax.set_title(title, fontproperties=zhfont, fontsize=14)
        ax.set_xlabel(xlabel, fontproperties=zhfont, fontsize=12)
        ax.set_ylabel(ylabel, fontproperties=zhfont, fontsize=12)
        ax.legend(prop=zhfont, fontsize=10)
    else:
        ax.set_title(title, fontsize=14)
        ax.set_xlabel(xlabel, fontsize=12)
        ax.set_ylabel(ylabel, fontsize=12)
        ax.legend(fontsize=10)


def _add_doc_table(docu, headers, rows):
    table = docu.add_table(rows=1, cols=len(headers), style='Table Grid')
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for row in rows:
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
    return table


def _phi_scatter_chart(workpath, green):
    """绿谱线 ±1 级各次测量衍射角散点图（+ 平均线）。"""
    zhfont = _font()
    plus = green["plus"]
    minus = green["minus"]
    fig, ax = plt.subplots(figsize=(8, 5))
    if plus:
        ax.plot(np.arange(1, len(plus) + 1), plus, 'o', color='#4472C4',
                markersize=7, label='绿(+1) 级', zorder=5)
    if minus:
        ax.plot(np.arange(1, len(minus) + 1) + 3.5, minus, 's', color='#E74C3C',
                markersize=7, label='绿(-1) 级', zorder=5)
    ax.axhline(green["mean"], color='#3A9D5D', linestyle='--', linewidth=1.5,
               label='平均值 {:.4f}°'.format(green["mean"]))
    _decorate(ax, "绿谱线衍射角六次测量（±1 级）", "测量序号", "φ / °")
    ax.grid(True, linestyle='--', alpha=0.4)
    fig.tight_layout()
    filename = "chart_phi.png"
    fig.savefig(os.path.join(workpath, filename), dpi=150, bbox_inches='tight')
    plt.close(fig)
    return filename


def _lambda_sin_chart(workpath, items):
    """λ-sinφ 关系图（验证光栅方程）。items: [(谱线名, sinφ, λ/nm, 理论值nm)]。

    过原点拟合 λ = d·sinφ，斜率即光栅常数，可与绿光法求得的 d 对照。
    返回 (文件名, 拟合斜率 d_fit)。
    """
    zhfont = _font()
    xs = np.array([it[1] for it in items], dtype=float)
    ys = np.array([it[2] for it in items], dtype=float)
    d_fit = float(np.sum(xs * ys) / np.sum(xs ** 2))
    x_fit = np.linspace(0.0, xs.max() * 1.12, 100)
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(xs, ys, 'o', color='#4472C4', markersize=7, zorder=5, label='测量点')
    ax.plot(x_fit, d_fit * x_fit, '-', color='#E74C3C', linewidth=2,
            label='过原点拟合 λ = {:.0f}·sinφ'.format(d_fit))
    for it in items:
        kwargs = dict(xy=(it[1], it[2]), textcoords="offset points",
                      xytext=(8, -6), fontsize=10)
        if zhfont:
            kwargs["fontproperties"] = zhfont
        ax.annotate(it[0], **kwargs)
    _decorate(ax, "λ-sinφ 关系（验证光栅方程）", "sinφ", "λ / nm")
    ax.grid(True, linestyle='--', alpha=0.4)
    fig.tight_layout()
    filename = "chart_lambda_sin.png"
    fig.savefig(os.path.join(workpath, filename), dpi=150, bbox_inches='tight')
    plt.close(fig)
    return filename, d_fit


def _lambda_bar_chart(workpath, items):
    """三条谱线"测量值 vs 理论值"柱状图。items: [(谱线名, 测量值nm, 理论值nm)]。"""
    zhfont = _font()
    names = [it[0] for it in items]
    calc = [it[1] for it in items]
    theory = [it[2] for it in items]
    x = np.arange(len(items))
    width = 0.35
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.bar(x - width / 2, calc, width, color='#4472C4', label='测量值', zorder=3)
    ax.bar(x + width / 2, theory, width, color='#A5A5A5', label='理论值', zorder=3)
    for i, (c, t) in enumerate(zip(calc, theory)):
        dev = (c - t) / t * 100.0
        ax.text(i, max(c, t), "{:+.2f}%".format(dev), ha='center', va='bottom', fontsize=10)
    lo, hi = min(min(calc), min(theory)), max(max(calc), max(theory))
    span = hi - lo if hi > lo else 1.0
    ax.set_ylim(lo - 0.15 * span, hi + 0.20 * span)
    ax.set_xticks(x)
    if zhfont:
        ax.set_xticklabels(names, fontproperties=zhfont)
        ax.set_title("各谱线波长测量值与理论值（k = 1）", fontproperties=zhfont, fontsize=14)
        ax.set_ylabel("波长 λ / nm", fontproperties=zhfont)
        ax.legend(prop=zhfont)
    else:
        ax.set_xticklabels(names)
        ax.set_title("各谱线波长测量值与理论值（k = 1）", fontsize=14)
        ax.set_ylabel("波长 λ / nm")
        ax.legend()
    ax.grid(True, axis='y', linestyle='--', alpha=0.4)
    fig.tight_layout()
    filename = "chart_lambda.png"
    fig.savefig(os.path.join(workpath, filename), dpi=150, bbox_inches='tight')
    plt.close(fig)
    return filename


def _mean(values):
    return float(sum(values) / len(values)) if values else None


# ─────────────────────────────────────────────
# 计算：衍射角 → 光栅常数 → 各谱线波长
# ─────────────────────────────────────────────
def _calc_grating(rows):
    """完整计算：绿谱线 ±1 级衍射角（分侧统计 + 对称性检验）、光栅常数、
    三条谱线波长与偏差。返回 {"lines", "warnings", "green", "others"}。"""
    zero_row = _find_zero_row(rows)
    if zero_row is None:
        raise ValueError("数据表缺少「0级位置」行（行名需包含“0”和“级”，如“0级位置”），无法计算衍射角。")

    warnings = []
    lines = []

    green_plus = _collect_phis_side(rows, "绿", "+")
    green_minus = _collect_phis_side(rows, "绿", "-")
    all_green = green_plus + green_minus
    if len(all_green) < 2:
        raise ValueError("绿谱线 ±1 级有效衍射角不足 2 个，无法计算光栅常数。")

    res_phi = analyse(pd.Series(all_green), _DELTA_INSTRUMENT, _DELTA_ESTIMATE,
                      "φ", "°", confidence_C=3 ** 0.5)
    # 角度以"度"代入，表达式内显式乘 pi/180（sympy 的三角函数按弧度计算）
    res_d = analyse_com("d=λ_g/sin(phi*pi/180)",
                        (("phi", float(res_phi.average), float(res_phi.unc)),),
                        (("λ_g", _LAMBDA_GREEN),), "nm")
    d_nm = float(res_d.ans)
    u_d_nm = float(res_d.unc)

    green = {"plus": green_plus, "minus": green_minus, "phis": all_green,
             "plus_mean": _mean(green_plus), "minus_mean": _mean(green_minus),
             "mean": float(res_phi.average), "unc": float(res_phi.unc),
             "res_phi": res_phi, "res_d": res_d,
             "d_nm": d_nm, "u_d_nm": u_d_nm, "diff_arcmin": None}

    # 指导书表 2.2：±1 级衍射角分别记录，再取平均
    if green_plus:
        lines.append("绿(+1) 级衍射角：{}（均值 {:.4f}°）".format(
            '、'.join('%.5g°' % v for v in green_plus), green["plus_mean"]))
    if green_minus:
        lines.append("绿(-1) 级衍射角：{}（均值 {:.4f}°）".format(
            '、'.join('%.5g°' % v for v in green_minus), green["minus_mean"]))
    if green_plus and green_minus:
        diff_arcmin = (green["plus_mean"] - green["minus_mean"]) * 60.0
        green["diff_arcmin"] = diff_arcmin
        if abs(diff_arcmin) <= _SYMMETRY_LIMIT_ARCMIN:
            lines.append("两侧衍射角相差 {:.1f}′（≤3′），对称性良好，光栅调节符合要求。".format(diff_arcmin))
        else:
            warnings.append("绿谱线两侧衍射角相差 {:.1f}′（>3′），可能光栅刻痕未与仪器转轴平行或未保证垂直入射，建议重新调节。".format(diff_arcmin))
    else:
        warnings.append("绿谱线 ±1 级数据不完整（+{} 个/−{} 个），未做两侧对称性检验。".format(
            len(green_plus), len(green_minus)))

    lines.append("绿谱线平均衍射角 φ ≈ {}°（U ≈ {}°，P=0.95）".format(
        formatted(res_phi.average, 3), formatted(res_phi.unc, 3)))
    lines.append("光栅常数 d ≈ {} nm ≈ {} μm，即空间频率约 {} 线/mm。".format(
        formatted(d_nm, 1), formatted(d_nm / 1000.0, 4), formatted(1e6 / d_nm, 1)))
    if d_nm < _D_MIN_NM or d_nm > _D_MAX_NM:
        warnings.append("光栅常数 d ≈ {} μm 与常见教学光栅（约 3.3~10 μm，100~300 线/mm）差异较大，请核对数据。".format(
            formatted(d_nm / 1000.0, 3)))

    others = {}
    for line_name in _LINE_ORDER:
        plus = _collect_phis_side(rows, line_name, "+")
        minus = _collect_phis_side(rows, line_name, "-")
        allp = plus + minus
        if len(allp) < 2:
            warnings.append("谱线「{}」±1 级有效衍射角不足 2 个，未计算其波长。".format(line_name))
            continue
        res_phi_l = analyse(pd.Series(allp), _DELTA_INSTRUMENT, _DELTA_ESTIMATE,
                            "φ", "°", confidence_C=3 ** 0.5)
        res_lambda = analyse_com("λ=d*sin(phi*pi/180)",
                                 (("d", d_nm, u_d_nm),
                                  ("phi", float(res_phi_l.average), float(res_phi_l.unc))),
                                 (), "nm")
        lam_nm = float(res_lambda.ans)
        theory_nm = _LINE_THEORY[line_name]
        dev = (lam_nm - theory_nm) / theory_nm * 100.0
        info = {"plus": plus, "minus": minus,
                "plus_mean": _mean(plus), "minus_mean": _mean(minus),
                "mean": float(res_phi_l.average), "unc": float(res_phi_l.unc),
                "res_phi": res_phi_l, "res_lambda": res_lambda,
                "lam_nm": lam_nm, "theory_nm": theory_nm, "dev": dev}
        others[line_name] = info
        lines.append("{}谱线：衍射角 φ ≈ {}°，λ ≈ {} nm（标准值 {} nm，偏差 {:+.2f}%）".format(
            line_name, formatted(info["mean"], 3), formatted(lam_nm, 2), theory_nm, dev))
        if abs(dev) > _DEV_WARN_PERCENT:
            warnings.append("{}谱线波长与标准值偏差 {:+.2f}%（>1%），请检查谱线识别与级次。".format(
                line_name, dev))
    if not others:
        warnings.append("蓝紫/黄内/黄外谱线均无有效数据，未计算波长。")

    return {"lines": lines, "warnings": warnings, "green": green, "others": others}


# ─────────────────────────────────────────────
# 新版结构化接口
# ─────────────────────────────────────────────
def schema():
    sample = load_sample_data("exp18_a", "分光计的调节与使用")
    table1 = make_table(
        "table1", "分光计光栅衍射角测量原始数据",
        _TABLE_LABELS,
        sample=sample,
        readonly=(7, 8, 9, 10),
        text_columns=(0,),
        min_rows=9,
        description="按指导书数据表顺序记录：第 1 行为「0级位置」（行名必须含“0”和“级”），"
                    "其后为绿/蓝紫/黄内/黄外各谱线的 ±1 级位置，共 9 行。"
                    "θ1/θ2 为左、右游标读数，各测 3 次。角度可输入 \"123°30′\" 或十进制度"
                    "（123.5 = 123°30′）。衍射角 φ = |θ_k − θ_0|（左右游标平均，消除偏心误差），"
                    "后 4 列由系统自动计算。",
    )
    table1["calc"] = {"label": "计算衍射角、光栅常数与波长"}

    return make_schema(
        (
            "分光计调节与光栅衍射：先完成望远镜、平行光管与光栅位置的调节（见实验指导），"
            "记录 0 级位置与各谱线 ±1 级的左右游标读数（各 3 次），"
            "由光栅方程 " r"$d\sin\varphi = k\lambda$" " 求光栅常数 " r"$d$"
            "（绿光 546.07 nm），再求汞灯蓝紫线与两条黄线的波长。"
        ),
        [table1],
        analysis_hints="绿谱线两侧（+1、-1 级）的衍射角应基本一致（相差 ≤3′），"
                       "若相差较大说明光栅刻痕未与仪器转轴平行或未保证垂直入射；"
                       "由 d = λ绿/sinφ 得光栅常数（教学光栅常见 100~300 线/mm，即 d 约 3.3~10 μm）；"
                       "再由 λ = d·sinφ 求蓝紫线与两条黄线的波长，应与标准值 435.83、576.96、579.07 nm 接近"
                       "（相对偏差一般 <1%）。",
        preview_enabled=True,
        report_enabled=False,
        formulas=get_formulas("exp18_a"),
        variables=get_variables("exp18_a"),
        table_theory=get_table_theory("exp18_a"),
    )


def _compute_all(payload):
    """补全每行衍射角列并计算光栅常数/波长（富结果，供 handle_structured 使用）。"""
    tables = copied_tables(payload)
    rows = tables.get("table1", [])
    zero_row = _find_zero_row(rows)
    for row in rows:
        info = _row_info(row.get("c0", ""))
        if info[0] != "line" or zero_row is None:
            for col in _PHI_COLS + (_PHI_MEAN_COL,):
                row[col] = ""
            continue
        phis = []
        for m, col in enumerate(_PHI_COLS):
            phi = _phi_measurement(row, zero_row, m)
            row[col] = formatted(phi, 3) if phi is not None else ""
            if phi is not None:
                phis.append(phi)
        row[_PHI_MEAN_COL] = formatted(sum(phis) / len(phis), 3) if phis else ""

    calc_results = {}
    calc_messages = []
    fit_notes = {}
    if any(_has_value(r.get(c)) for r in rows for c in ("c1", "c2", "c3", "c4", "c5", "c6")):
        try:
            calc_results["table1"] = _calc_grating(rows)
            fit_notes["table1"] = [
                "光栅方程（垂直入射）：d·sinφ = kλ（k = 0, ±1, ±2, …）。",
                "衍射角 φ = |θ_k − θ_0|，左、右游标平均以消除刻度盘偏心误差。",
                "绿谱线两侧（+1、-1 级）衍射角应基本一致（相差 ≤3′），不一致说明光栅未调好。",
                "d = λ绿/sinφ（λ绿 = 546.07 nm）；λ = d·sinφ。",
            ]
        except ValueError as exc:
            calc_messages.append(str(exc))
    return {"tables": tables, "calc_results": calc_results,
            "calc_messages": calc_messages, "fit_notes": fit_notes}


def preview(payload):
    """实时预计算接口：只返回前端需要且可 JSON 序列化的结论行。"""
    raw = _compute_all(payload)
    calc_results = {}
    for tid, v in raw["calc_results"].items():
        lines = list(v["lines"])
        lines.extend("注意：" + w for w in v["warnings"])
        calc_results[tid] = {"lines": lines}
    return {"tables": raw["tables"], "calc_results": calc_results,
            "calc_messages": raw["calc_messages"], "fit_notes": raw["fit_notes"]}


def handle_structured(workpath, payload):
    raw = _compute_all(payload)
    rows = raw["tables"].get("table1", [])
    calc = raw["calc_results"].get("table1")

    # 数据行（谱线名或任一读数列有内容）
    data_rows = []
    for r in rows:
        if not any(str(r.get(c, "")).strip() for c in
                   ("c0",) + tuple(col for pair in _MEASUREMENT_COLS for col in pair)):
            continue
        data_rows.append([str(r.get(c, "")) for c in
                          ("c0", "c1", "c2", "c3", "c4", "c5", "c6",
                           "c7", "c8", "c9", "c10")])
    if not data_rows:
        return {"code": 1, "message": "请先填写测量数据表：至少需要「0级位置」与绿谱线 ±1 级的游标读数。"}

    warnings = list(raw["calc_messages"])
    if calc:
        warnings.extend(calc["warnings"])
    summary = []
    charts = []

    docu = Document()
    style_doc_font(docu)
    docu.add_heading(name(), level=0)

    # ── 一、测量原始数据 ──────────────────────────────────────────────
    docu.add_heading("一、光栅衍射角测量原始数据", level=1)
    _add_doc_table(docu, _TABLE_LABELS, data_rows)
    docu.add_paragraph("θ1/θ2 为左、右游标读数；衍射角 φ = |θ_k − θ_0|（左右游标平均，消除偏心误差），"
                       "后 4 列由系统自动计算。")
    docu.add_paragraph()

    if calc is None:
        docu.add_paragraph("（缺少有效的绿谱线数据，未计算光栅常数与波长。）")
        docu.save(os.path.join(workpath, name() + ".docx"))
        return {"code": 0,
                "summary": summary or ["分光计数据已处理，详见 Word 文档。"],
                "warnings": warnings, "charts": charts,
                "document": name() + ".docx"}

    green = calc["green"]
    others = calc["others"]

    # ── 二、绿谱线衍射角与光栅常数 ─────────────────────────────────────
    docu.add_heading("二、绿谱线衍射角与光栅常数", level=1)
    if green["plus"]:
        docu.add_paragraph("绿(+1) 级衍射角（3 次测量，单位 °）：" +
                           ', '.join(['%.5g' % v for v in green["plus"]]))
    if green["minus"]:
        docu.add_paragraph("绿(-1) 级衍射角（3 次测量，单位 °）：" +
                           ', '.join(['%.5g' % v for v in green["minus"]]))
    if green["diff_arcmin"] is not None:
        docu.add_paragraph("两侧衍射角相差 {:.1f}′（≤3′ 认为一致）。".format(green["diff_arcmin"]))
    insert_data(docu, "绿谱线衍射角 φ（±1 级共 {} 次测量）".format(len(green["phis"])),
                green["res_phi"], "word")
    docu.add_paragraph("将绿光波长 λ_g = 546.07 nm 与平均衍射角代入光栅方程 d·sinφ = kλ（k = 1），得光栅常数")
    insert_data_com(docu, "光栅常数 d = λ_g/sinφ", green["res_d"], "word")
    docu.add_paragraph("光栅常数 d ≈ {} nm ≈ {} μm，即空间频率约 {} 线/mm。".format(
        formatted(green["d_nm"], 1), formatted(green["d_nm"] / 1000.0, 4),
        formatted(1e6 / green["d_nm"], 1)))
    fn = _phi_scatter_chart(workpath, green)
    charts.append({"filename": fn, "title": "绿谱线衍射角六次测量"})
    docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
    summary.append("绿谱线平均衍射角 φ ≈ {}°（U ≈ {}°，P=0.95）".format(
        formatted(green["mean"], 3), formatted(green["unc"], 3)))
    summary.append("光栅常数 d ≈ {} nm（≈ {} μm，约 {} 线/mm）".format(
        formatted(green["d_nm"], 1), formatted(green["d_nm"] / 1000.0, 4),
        formatted(1e6 / green["d_nm"], 1)))
    docu.add_paragraph()

    # ── 三、汞灯其他谱线的波长 ─────────────────────────────────────────
    docu.add_heading("三、汞灯其他谱线的波长", level=1)
    result_items = []
    sin_items = [("绿", float(np.sin(np.radians(green["mean"]))), _LAMBDA_GREEN, _LAMBDA_GREEN)]
    for line_name in _LINE_ORDER:
        info = others.get(line_name)
        if info is None:
            continue
        docu.add_paragraph("{}谱线：φ(+1) 均值 {:.4f}°、φ(-1) 均值 {:.4f}°，"
                           "平均衍射角 {:.4f}°".format(
                               line_name,
                               info["plus_mean"] if info["plus_mean"] is not None else float('nan'),
                               info["minus_mean"] if info["minus_mean"] is not None else float('nan'),
                               info["mean"]))
        insert_data_com(docu, "{}谱线波长 λ".format(line_name), info["res_lambda"], "word")
        docu.add_paragraph("{}谱线实测值与标准值 {} nm 的相对偏差为 {:+.2f}%。".format(
            line_name, info["theory_nm"], info["dev"]))
        docu.add_paragraph()
        result_items.append((line_name, info["lam_nm"], info["theory_nm"]))
        sin_items.append((line_name, float(np.sin(np.radians(info["mean"]))),
                          info["lam_nm"], info["theory_nm"]))
        summary.append("{}谱线波长 λ ≈ {} nm（标准值 {} nm，相对偏差 {:+.2f}%）".format(
            line_name, formatted(info["lam_nm"], 2), info["theory_nm"], info["dev"]))

    if len(sin_items) >= 2:
        fn, d_fit = _lambda_sin_chart(workpath, sin_items)
        charts.append({"filename": fn, "title": "λ-sinφ 关系（验证光栅方程）"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))
        dev_d_fit = (d_fit - green["d_nm"]) / green["d_nm"] * 100.0
        docu.add_paragraph("λ-sinφ 过原点拟合斜率 d_fit ≈ {} nm，与绿光法求得的 d 相对偏差 {:+.2f}%。".format(
            formatted(d_fit, 0), dev_d_fit))
        summary.append("λ-sinφ 过原点拟合斜率 d_fit ≈ {} nm".format(formatted(d_fit, 0)))
        if abs(dev_d_fit) > 10.0:
            warnings.append("λ-sinφ 拟合斜率与绿光法光栅常数偏差 {:.1f}%（>10%），请检查各谱线衍射角数据。".format(dev_d_fit))

    if result_items:
        fn = _lambda_bar_chart(workpath, result_items)
        charts.append({"filename": fn, "title": "各谱线波长测量值与理论值"})
        docu.add_picture(os.path.join(workpath, fn), width=Inches(6))

    docu.save(os.path.join(workpath, name() + ".docx"))

    return {
        "code": 0,
        "summary": summary or ["分光计数据已处理，详见 Word 文档。"],
        "warnings": warnings,
        "charts": charts,
        "document": name() + ".docx",
    }


# ─────────────────────────────────────────────
# 旧版接口（向后兼容）
# ─────────────────────────────────────────────
def _rows_from_legacy_frame(data):
    """把旧版 CSV 读出的 DataFrame 转成结构化行对象（c0~c6）。"""
    rows = []
    for _, record in data.iterrows():
        row = {"c0": str(record["line"]).strip() if pd.notna(record["line"]) else ""}
        for col in ("c1", "c2", "c3", "c4", "c5", "c6"):
            value = record[col]
            row[col] = str(value).strip() if pd.notna(value) else ""
        rows.append(row)
    return rows


def handle(workpath, extension):
    """旧版 CSV 单表接口（兼容 b_adapter 等旧调用路径）：
    读取谱线表原始数据（谱线 + 3 次测量的 θ1/θ2）→ 光栅常数 d 与各谱线波长 → Word。"""
    try:
        excelpath = workpath + name() + '.' + extension
        column_names = ["line", "c1", "c2", "c3", "c4", "c5", "c6"]

        if extension == 'csv':
            with open(excelpath, 'rb') as f:
                encode = chardet.detect(f.read())["encoding"]
            data = pd.read_csv(excelpath, header=0, names=column_names, encoding=encode)
        else:
            data = pd.read_excel(excelpath, header=0, names=column_names)

        os.remove(excelpath)

        payload = {"parameters": {}, "tables": {"table1": _rows_from_legacy_frame(data)}}
        result = handle_structured(workpath, payload)
        return 0 if result.get("code") == 0 else 1
    except Exception:
        traceback.print_exc()  # 打印错误
        return 1  # 若失败，返回1
