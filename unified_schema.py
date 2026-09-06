"""统一实验 Schema 与新版插件适配工具。

这个模块只负责框架层的转换：把插件的 ``table_configs`` 变成通用
Schema，并把浏览器提交的结构化表格转换回插件能够处理的二维数组。
具体物理公式仍由各实验插件负责。
"""

from __future__ import annotations

import html
import os
import re
from typing import Any

from experiment_schema_overrides import apply_schema_overrides


def _style_doc_font(document) -> None:
    """统一 Word 文档字体为微软雅黑、加粗。

    覆盖 Normal 正文样式与 Title / Heading 1~3 标题样式，使标题、正文、
    表格表头与单元格全部使用微软雅黑加粗。Heading 默认走主题字体（中文宋体、
    西文 Calibri），不显式覆盖会与正文不一致，故在此一并处理。
    """
    from docx.oxml.ns import qn

    font_name = "微软雅黑"
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.bold = True
    rpr = normal._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = font_name
        style.font.bold = True
        s_rpr = style._element.get_or_add_rPr()
        s_rfonts = s_rpr.find(qn("w:rFonts"))
        if s_rfonts is None:
            s_rfonts = s_rpr.makeelement(qn("w:rFonts"), {})
            s_rpr.append(s_rfonts)
        s_rfonts.set(qn("w:eastAsia"), font_name)


def html_to_text(value: Any) -> str:
    """把表头中的简单 HTML 上下标转换为可读纯文本。"""
    text = str(value or "")
    text = re.sub(r"<\s*sub\s*>(.*?)<\s*/\s*sub\s*>", r"_(\1)", text, flags=re.I)
    text = re.sub(r"<\s*sup\s*>(.*?)<\s*/\s*sup\s*>", r"^(\1)", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    return html.unescape(text).strip()


def _column_id(index: int) -> str:
    return f"c{index}"


def _plugin_object(plugin: Any) -> Any:
    """注册中心既可能保存类，也可能保存适配器实例。"""
    return plugin() if isinstance(plugin, type) else plugin


def schema_from_plugin(plugin: Any, experiment_id: str) -> dict[str, Any] | None:
    """把新版插件的 ``table_configs`` 转换成统一 Schema。"""
    table_configs = getattr(plugin, "table_configs", None)
    ui_schema = getattr(plugin, "ui_schema", None)
    preview_enabled = callable(getattr(plugin, "preview", None))

    if callable(ui_schema):
        schema = ui_schema()
        schema.setdefault("schema_version", 2)
        schema.setdefault("schema_revision", 1)
        schema.setdefault("draft_enabled", True)
        schema.setdefault("preview_enabled", preview_enabled)
        return apply_schema_overrides(experiment_id, schema)
    if isinstance(ui_schema, dict):
        schema = dict(ui_schema)
        schema.setdefault("schema_version", 2)
        schema.setdefault("schema_revision", 1)
        schema.setdefault("draft_enabled", True)
        schema.setdefault("preview_enabled", preview_enabled)
        return apply_schema_overrides(experiment_id, schema)
    if not table_configs:
        return None

    tables = []
    for table_config in table_configs:
        columns = []
        for index, raw_column in enumerate(table_config.get("columns", [])):
            if isinstance(raw_column, dict):
                column = dict(raw_column)
                column.setdefault("id", _column_id(index))
                column["label"] = html_to_text(column.get("label", column["id"]))
            else:
                column = {
                    "id": _column_id(index),
                    "label": html_to_text(raw_column),
                    "type": "text",
                }
            columns.append(column)

        table = {
            "id": table_config.get("id", f"table{len(tables) + 1}"),
            "title": html_to_text(table_config.get("title", "数据表")),
            "description": html_to_text(table_config.get("description", "")),
            "required": table_config.get("required", True),
            "min_rows": table_config.get("min_rows", 1),
            "initial_rows": table_config.get("row_count", table_config.get("initial_rows", 3)),
            "columns": columns,
            "sample": table_config.get("sample", []),
        }

        chart = table_config.get("chart")
        if chart:
            table["chart"] = dict(chart)
        elif table_config.get("x_label") and table_config.get("y_label") and len(columns) >= 2:
            table["chart"] = {
                "x_column": table_config.get("x_column", columns[0]["id"]),
                "y_column": table_config.get("y_column", columns[1]["id"]),
                "x_label": html_to_text(table_config["x_label"]),
                "y_label": html_to_text(table_config["y_label"]),
                "title": html_to_text(table_config.get("chart_title", table["title"])),
                "fit": table_config.get("fit", "linear"),
            }
        tables.append(table)

    return apply_schema_overrides(experiment_id, {
        "schema_version": 2,
        "schema_revision": getattr(plugin, "schema_revision", 1),
        "draft_enabled": True,
        "description": html_to_text(getattr(plugin, "description", "")),
        "parameters": list(getattr(plugin, "parameters", [])),
        "tables": tables,
        "analysis_hints": getattr(plugin, "analysis_hints", ""),
        "preview_enabled": preview_enabled,
        "report_enabled": True,
    })


def _ordered_rows(schema: dict[str, Any], payload: dict[str, Any]) -> dict[str, list[list[str]]]:
    submitted_tables = payload.get("tables") or {}
    result: dict[str, list[list[str]]] = {}
    for table in schema.get("tables", []):
        table_id = table.get("id")
        column_ids = [column.get("id") for column in table.get("columns", [])]
        rows = []
        for raw_row in submitted_tables.get(table_id, []) or []:
            rows.append([str(raw_row.get(column_id, "")) for column_id in column_ids])
        result[table_id] = rows
    return result


def _plugin_constants(schema: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    submitted = payload.get("parameters") or {}
    constants: dict[str, Any] = {}
    for parameter in schema.get("parameters", []):
        parameter_id = parameter.get("id")
        value = submitted.get(parameter_id, parameter.get("default", ""))
        constants[parameter_id] = value
        backend_key = parameter.get("backend_key")
        if backend_key:
            constants[backend_key] = value
    return constants


def _flatten_results(results: Any) -> tuple[list[str], list[str]]:
    if not isinstance(results, dict):
        return [str(results)], []

    summary: list[str] = []
    warnings: list[str] = []
    for section_name in ("steps", "final"):
        section = results.get(section_name)
        if isinstance(section, dict):
            for key, value in section.items():
                summary.append(f"{key}：{value}")
    message = results.get("message")
    if message:
        summary.append(str(message))
    raw_warnings = results.get("warnings")
    if isinstance(raw_warnings, list):
        warnings.extend(str(item) for item in raw_warnings)
    if not summary:
        summary.append("数据已接收，可继续进行异常检测、图表生成或实验报告生成。")
    return summary, warnings


def _write_docx(
    workpath: str,
    title: str,
    schema: dict[str, Any],
    payload: dict[str, Any],
    summary: list[str],
) -> str:
    from docx import Document

    document = Document()
    _style_doc_font(document)  # 统一标题/正文/表头为微软雅黑加粗
    document.add_heading(title, level=0)

    parameters = payload.get("parameters") or {}
    if parameters:
        document.add_heading("实验参数", level=1)
        for parameter in schema.get("parameters", []):
            parameter_id = parameter.get("id")
            if parameter_id in parameters:
                label = parameter.get("label", parameter_id)
                unit = parameter.get("unit", "")
                document.add_paragraph(f"{label}：{parameters[parameter_id]} {unit}".strip())

    submitted_tables = payload.get("tables") or {}
    for table_schema in schema.get("tables", []):
        table_id = table_schema.get("id")
        rows = submitted_tables.get(table_id, []) or []
        if not rows:
            continue
        document.add_heading(table_schema.get("title", table_id), level=1)
        columns = table_schema.get("columns", [])
        table = document.add_table(rows=1, cols=len(columns), style="Table Grid")
        for index, column in enumerate(columns):
            cell = table.cell(0, index)
            cell.text = str(column.get("label", column.get("id", "")))
            for run in cell.paragraphs[0].runs:
                run.font.bold = True  # 表头加粗
        for raw_row in rows:
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                cells[index].text = str(raw_row.get(column.get("id"), ""))

    document.add_heading("计算与处理结果", level=1)
    for item in summary:
        document.add_paragraph(item)

    safe_title = re.sub(r'[\\/:*?"<>|]+', "_", title).strip() or "实验数据"
    filename = f"{safe_title}.docx"
    document.save(os.path.join(workpath, filename))
    return filename


def handle_plugin_structured(
    plugin: Any,
    experiment_id: str,
    workpath: str,
    payload: dict[str, Any],
) -> dict[str, Any]:
    """用统一协议调用新版插件，并生成可下载的基础 Word 文档。"""
    schema = schema_from_plugin(plugin, experiment_id)
    if not schema:
        return {"code": 1, "message": "该实验没有可用的结构化表格配置"}

    plugin_object = _plugin_object(plugin)
    tables = _ordered_rows(schema, payload)
    constants = _plugin_constants(schema, payload)
    try:
        results = plugin_object.calculate(tables, constants)
    except Exception as exc:
        return {"code": 1, "message": f"实验计算失败：{exc}"}

    if isinstance(results, dict) and results.get("status") == "error":
        return {"code": 1, "message": results.get("message", "实验计算失败")}

    summary, warnings = _flatten_results(results)
    filename = _write_docx(
        workpath,
        html_to_text(getattr(plugin, "name", experiment_id)),
        schema,
        payload,
        summary,
    )
    return {
        "code": 0,
        "summary": summary,
        "warnings": warnings,
        "charts": [],
        "document": filename,
    }
